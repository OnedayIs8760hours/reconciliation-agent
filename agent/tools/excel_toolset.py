"""面向 LLM 的 Excel 操作工具集。

这些工具有意暴露小型、类型化、确定性的操作，而不是
允许 LLM 直接修改工作簿。工作流可将它们用作
需批准的操作：先检查，再写入明确的输出路径。
"""

from copy import copy
from pathlib import Path
from typing import Any

from openpyxl.styles import PatternFill
from openpyxl.utils import range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

from agent.core.formula_manager import FormulaManager
from agent.core.view_manager import ViewManager
from agent.core.workbook_loader import WorkbookLoader
from agent.tools.excel_models import (
    ExcelOperationName,
    ExcelSafetyPolicy,
    ExcelToolObservation,
    ExcelToolRequest,
    RowMark,
)


class ExcelToolError(RuntimeError):
    """当面向 LLM 的 Excel 操作违反工具约束时抛出。"""


class ExcelToolset:
    """可以安全暴露给 LLM 智能体的确定性 Excel 能力。"""

    MUTATING_OPERATIONS = {
        ExcelOperationName.WRITE_CELL,
        ExcelOperationName.WRITE_ROW,
        ExcelOperationName.APPEND_COLUMNS,
        ExcelOperationName.INSERT_ROWS,
        ExcelOperationName.COPY_ROW_STYLE,
        ExcelOperationName.MARK_ROWS,
        ExcelOperationName.SET_FORMULA,
        ExcelOperationName.RESET_VIEW,
        ExcelOperationName.SAVE_AS,
    }

    def __init__(
        self,
        loader: WorkbookLoader | None = None,
        policy: ExcelSafetyPolicy | None = None,
    ) -> None:
        self.loader = loader or WorkbookLoader()
        self.policy = policy or ExcelSafetyPolicy()
        self.formula_manager = FormulaManager()
        self.view_manager = ViewManager()

    def execute(self, request: ExcelToolRequest) -> ExcelToolObservation:
        """执行一个类型化 Excel 操作并返回有界观察结果。"""
        self._validate_request(request)
        if request.operation == ExcelOperationName.INSPECT_WORKBOOK:
            return self.inspect_workbook(request)
        if request.operation == ExcelOperationName.READ_RANGE:
            return self.read_range(request)

        workbook = self.loader.load(request.workbook_path)
        sheet = self._sheet(workbook, request.sheet_name)
        changed_cells: list[str] = []
        changed_rows: list[int] = []

        if request.operation == ExcelOperationName.WRITE_CELL:
            changed_cells = self._write_cells(sheet, request)
        elif request.operation == ExcelOperationName.WRITE_ROW:
            changed_cells, changed_rows = self._write_rows(sheet, request)
        elif request.operation == ExcelOperationName.APPEND_COLUMNS:
            changed_cells = self._append_columns(sheet, request.append_headers)
        elif request.operation == ExcelOperationName.INSERT_ROWS:
            changed_rows = self._insert_rows(sheet, request)
        elif request.operation == ExcelOperationName.COPY_ROW_STYLE:
            changed_cells = self._copy_row_style(sheet, request)
        elif request.operation == ExcelOperationName.MARK_ROWS:
            changed_cells, changed_rows = self._mark_rows(sheet, request.row_marks)
        elif request.operation == ExcelOperationName.SET_FORMULA:
            changed_cells = self._set_formulas(sheet, request)
        elif request.operation == ExcelOperationName.RESET_VIEW:
            self.view_manager.reset(sheet, freeze_below_header=request.freeze_panes is not None)
            if request.freeze_panes:
                sheet.freeze_panes = request.freeze_panes
        elif request.operation == ExcelOperationName.SAVE_AS:
            pass
        else:
            raise ExcelToolError(f"Unsupported operation: {request.operation}")

        formula_errors = self.formula_manager.formula_error_cells(sheet)
        hash_risks = self.formula_manager.hash_display_risk_cells(sheet)
        warnings = []
        if formula_errors:
            warnings.append(f"发现公式错误单元格：{', '.join(formula_errors[:20])}")
        if hash_risks:
            warnings.append(f"发现金额显示########风险：{', '.join(hash_risks[:20])}")

        output_path = request.output_path or request.workbook_path
        if not request.dry_run:
            self.loader.save(workbook, output_path)
        return ExcelToolObservation(
            operation=request.operation,
            ok=True,
            workbook_path=request.workbook_path,
            output_path=output_path,
            sheet_name=sheet.title,
            message="dry-run完成，未写入文件" if request.dry_run else "Excel操作完成",
            changed_cells=changed_cells,
            changed_rows=changed_rows,
            headers=self._headers(sheet),
            preview=self._preview(sheet, max_rows=min(request.max_preview_rows, 10)),
            warnings=warnings,
        )

    def inspect_workbook(self, request: ExcelToolRequest) -> ExcelToolObservation:
        workbook = self.loader.load(request.workbook_path, data_only=False)
        sheet = self._sheet(workbook, request.sheet_name)
        return ExcelToolObservation(
            operation=ExcelOperationName.INSPECT_WORKBOOK,
            ok=True,
            workbook_path=request.workbook_path,
            sheet_name=sheet.title,
            message=f"工作簿包含{len(workbook.sheetnames)}个工作表；当前表{sheet.title}，{sheet.max_row}行 x {sheet.max_column}列",
            headers=self._headers(sheet),
            preview=self._preview(sheet, max_rows=request.max_preview_rows),
            warnings=self._inspection_warnings(sheet),
        )

    def read_range(self, request: ExcelToolRequest) -> ExcelToolObservation:
        workbook = self.loader.load(request.workbook_path, data_only=False)
        sheet = self._sheet(workbook, request.sheet_name)
        if not request.range_address:
            raise ExcelToolError("read_range requires range_address")
        min_col, min_row, max_col, max_row = range_boundaries(request.range_address)
        cell_count = (max_col - min_col + 1) * (max_row - min_row + 1)
        if cell_count > self.policy.max_cells_per_read:
            raise ExcelToolError(f"read_range exceeds max_cells_per_read={self.policy.max_cells_per_read}")
        preview = [
            [sheet.cell(row, col).value for col in range(min_col, max_col + 1)]
            for row in range(min_row, max_row + 1)
        ]
        return ExcelToolObservation(
            operation=ExcelOperationName.READ_RANGE,
            ok=True,
            workbook_path=request.workbook_path,
            sheet_name=sheet.title,
            message=f"已读取范围 {request.range_address}",
            preview=preview,
            headers=self._headers(sheet),
        )

    def _validate_request(self, request: ExcelToolRequest) -> None:
        if request.workbook_path.suffix.lower() not in self.policy.allowed_extensions:
            raise ExcelToolError(f"Unsupported workbook extension: {request.workbook_path.suffix}")
        if request.operation in self.MUTATING_OPERATIONS and self.policy.require_output_path_for_mutation:
            if request.output_path is None and self.policy.default_mutation_mode == "copy":
                raise ExcelToolError("Mutating Excel operations require output_path")
        write_count = len(request.cell_writes) + sum(len(row.values) for row in request.row_writes)
        write_count += len(request.formula_writes)
        if write_count > self.policy.max_cells_per_write:
            raise ExcelToolError(f"write request exceeds max_cells_per_write={self.policy.max_cells_per_write}")

    def _sheet(self, workbook: Any, sheet_name: str | None) -> Worksheet:
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise ExcelToolError(f"Sheet not found: {sheet_name}")
            return workbook[sheet_name]
        return workbook.active

    def _write_cells(self, sheet: Worksheet, request: ExcelToolRequest) -> list[str]:
        changed: list[str] = []
        for write in request.cell_writes:
            cell = sheet[write.coordinate]
            cell.value = write.value
            if write.number_format:
                cell.number_format = write.number_format
            changed.append(cell.coordinate)
        return changed

    def _write_rows(self, sheet: Worksheet, request: ExcelToolRequest) -> tuple[list[str], list[int]]:
        changed_cells: list[str] = []
        changed_rows: list[int] = []
        for row_write in request.row_writes:
            changed_rows.append(row_write.row_number)
            for column, value in row_write.values.items():
                cell = sheet.cell(row_write.row_number, int(column), value)
                changed_cells.append(cell.coordinate)
        return changed_cells, changed_rows

    def _append_columns(self, sheet: Worksheet, headers: list[str]) -> list[str]:
        existing = set(self._headers(sheet))
        changed: list[str] = []
        column = sheet.max_column + 1
        for header in headers:
            if header in existing:
                continue
            cell = sheet.cell(1, column, header)
            if column > 1:
                source = sheet.cell(1, column - 1)
                cell.font = copy(source.font)
                cell.fill = copy(source.fill)
                cell.border = copy(source.border)
                cell.alignment = copy(source.alignment)
            changed.append(cell.coordinate)
            column += 1
        return changed

    def _insert_rows(self, sheet: Worksheet, request: ExcelToolRequest) -> list[int]:
        if not request.insert_at:
            raise ExcelToolError("insert_rows requires insert_at")
        sheet.insert_rows(request.insert_at, request.insert_count)
        return list(range(request.insert_at, request.insert_at + request.insert_count))

    def _copy_row_style(self, sheet: Worksheet, request: ExcelToolRequest) -> list[str]:
        if not request.source_row or not request.target_row:
            raise ExcelToolError("copy_row_style requires source_row and target_row")
        changed: list[str] = []
        for col_idx in range(1, sheet.max_column + 1):
            source = sheet.cell(request.source_row, col_idx)
            target = sheet.cell(request.target_row, col_idx)
            target.font = copy(source.font)
            target.fill = copy(source.fill)
            target.border = copy(source.border)
            target.alignment = copy(source.alignment)
            target.number_format = source.number_format
            changed.append(target.coordinate)
        return changed

    def _mark_rows(self, sheet: Worksheet, marks: list[RowMark]) -> tuple[list[str], list[int]]:
        changed_cells: list[str] = []
        changed_rows: list[int] = []
        for mark in marks:
            status_col = self._ensure_header(sheet, mark.status_column_header)
            note_col = self._ensure_header(sheet, mark.note_column_header)
            fill = PatternFill(fill_type="solid", fgColor=mark.fill_color)
            status_cell = sheet.cell(mark.row_number, status_col, mark.status)
            note_cell = sheet.cell(mark.row_number, note_col, mark.note)
            changed_cells.extend([status_cell.coordinate, note_cell.coordinate])
            changed_rows.append(mark.row_number)
            for col_idx in range(1, sheet.max_column + 1):
                sheet.cell(mark.row_number, col_idx).fill = fill
        return changed_cells, changed_rows

    def _set_formulas(self, sheet: Worksheet, request: ExcelToolRequest) -> list[str]:
        changed: list[str] = []
        for formula_write in request.formula_writes:
            if not formula_write.formula.startswith("="):
                raise ExcelToolError(f"Formula must start with '=': {formula_write.coordinate}")
            cell = sheet[formula_write.coordinate]
            cell.value = formula_write.formula
            if formula_write.number_format:
                cell.number_format = formula_write.number_format
            changed.append(cell.coordinate)
        return changed

    def _ensure_header(self, sheet: Worksheet, header: str) -> int:
        for cell in sheet[1]:
            if cell.value == header:
                return cell.column
        column = sheet.max_column + 1
        sheet.cell(1, column, header)
        return column

    def _headers(self, sheet: Worksheet) -> list[str]:
        return [str(cell.value) for cell in sheet[1] if cell.value not in (None, "")]

    def _preview(self, sheet: Worksheet, *, max_rows: int) -> list[list[Any]]:
        rows: list[list[Any]] = []
        for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, max_rows), values_only=True):
            rows.append(list(row))
        return rows

    def _inspection_warnings(self, sheet: Worksheet) -> list[str]:
        warnings: list[str] = []
        formula_errors = self.formula_manager.formula_error_cells(sheet)
        hash_risks = self.formula_manager.hash_display_risk_cells(sheet)
        if formula_errors:
            warnings.append(f"发现公式错误单元格：{', '.join(formula_errors[:20])}")
        if hash_risks:
            warnings.append(f"发现金额显示########风险：{', '.join(hash_risks[:20])}")
        if sheet.merged_cells.ranges:
            warnings.append(f"发现{len(sheet.merged_cells.ranges)}个合并单元格，制表前应展开填充")
        return warnings
