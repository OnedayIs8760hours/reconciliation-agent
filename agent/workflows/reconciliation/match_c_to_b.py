"""Match the C workbook draft against B workbook rows using metadata mappings."""

from __future__ import annotations

import argparse
import copy
import json
import re
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from agent.workflows.reconciliation.copy_a_to_c import (
    get_dict,
    get_or_create_list,
    load_metadata,
    parse_numeric_month_day,
    require_string,
)
from tools import excel_tool

C_DATE_COL = 1
C_SPEC_COL = 2
C_QTY_COL = 3
C_PRICE_COL = 4
C_AMOUNT_COL = 5

APPEND_HEADERS = [
    "匹配状态",
    "匹配依据",
    "B表行号",
    "匹配方向",
    "差异金额",
    "仓库",
    "系统出入库时间",
    "单据编号",
    "货品编号",
    "货品名称",
    "规格",
    "入库数量",
    "入库成本单价",
    "入库成本金额",
    "出库数量",
    "出库成本单价",
    "出库成本金额",
]

B_TRACE_HEADERS = [
    "仓库",
    "系统出入库时间",
    "单据编号",
    "货品编号",
    "货品名称",
    "规格",
    "入库数量",
    "入库成本单价",
    "入库成本金额",
    "出库数量",
    "出库成本单价",
    "出库成本金额",
]

YELLOW_FILL = PatternFill("solid", fgColor="FFF2CC")
GREEN_FILL = PatternFill("solid", fgColor="E2F0D9")


@dataclass(frozen=True)
class BRecord:
    row_number: int
    date_key: tuple[int, int] | None
    spec: str
    values: dict[str, Any]

    def quantity_for_direction(self, direction: str) -> Decimal | None:
        header = "入库数量" if direction == "入库" else "出库数量"
        return to_decimal(self.values.get(header))


@dataclass(frozen=True)
class MatchPlan:
    status: str
    basis: str
    direction: str
    records: list[BRecord]
    amount_diff: Decimal | None = None


@dataclass(frozen=True)
class CColumns:
    date_col: int
    spec_col: int
    qty_col: int
    price_col: int
    amount_col: int
    first_detail_row: int


def match_c_to_b_from_metadata(
    metadata_path: Path,
    *,
    update_metadata: bool = True,
    force: bool = False,
) -> Path:
    """Run step 02: match C rows to B rows and append B trace fields."""

    metadata_path = metadata_path.resolve()
    metadata = load_metadata(metadata_path)
    c_path = Path(require_string(metadata, "c_file_path"))
    b_path = Path(require_string(metadata, "b_file_path"))

    if not force and has_done_step(metadata, "02_match_c_to_b"):
        return c_path

    summary = match_c_to_b(
        c_path,
        b_path,
        metadata,
    )

    if update_metadata:
        record_match_metadata(metadata_path, metadata, c_path=c_path, summary=summary)

    return c_path


def match_c_to_b(
    c_path: Path,
    b_path: Path,
    metadata: dict[str, Any],
) -> dict[str, int]:
    """Mutate the C workbook in place by appending B match details."""

    mapping_by_a_value = build_mapping_by_a_value(metadata)
    b_workbook = load_workbook(b_path, data_only=True)
    b_sheet = b_workbook.active
    b_records = read_b_records(b_sheet)

    c_workbook = load_workbook(c_path)
    c_sheet = c_workbook.active
    trim_trailing_empty_columns(c_sheet)
    c_columns = infer_c_columns(c_sheet)
    start_col = ensure_append_headers(c_sheet, header_row=c_columns.first_detail_row - 1)
    detail_rows = find_c_detail_rows(c_sheet, c_columns)
    cached_values = load_cached_a_values(metadata, c_columns)

    materialize_amount_formulas(c_sheet, detail_rows, c_columns, cached_values)
    plans = build_match_plans(c_sheet, detail_rows, mapping_by_a_value, b_records, c_columns)
    unmatched_b_records = find_unmatched_b_records(b_records, plans)
    apply_match_plans(c_sheet, plans, start_col=start_col)
    append_unmatched_b_records(c_sheet, unmatched_b_records, c_columns, start_col=start_col)
    rewrite_total_formula(c_sheet, c_columns)
    reset_sheet_view(c_sheet)

    c_workbook.save(c_path)
    summary = summarize_plans(plans)
    summary["appended_b_rows"] = len(unmatched_b_records)
    return summary


def build_mapping_by_a_value(metadata: dict[str, Any]) -> dict[str, list[str]]:
    product_mapping = get_dict(metadata, "product_mapping")
    result = get_dict(product_mapping, "result")
    mappings = result.get("mappings")
    if not isinstance(mappings, list):
        return {}

    mapping_by_a_value: dict[str, list[str]] = {}
    for item in mappings:
        if not isinstance(item, dict):
            continue
        a_value = clean_spec(item.get("a_value"))
        b_value = clean_spec(item.get("b_value"))
        if a_value and b_value:
            mapping_by_a_value.setdefault(a_value, [])
            if b_value not in mapping_by_a_value[a_value]:
                mapping_by_a_value[a_value].append(b_value)
    return mapping_by_a_value


def read_b_records(worksheet: Worksheet) -> list[BRecord]:
    header_map = read_header_map(worksheet)
    missing = [header for header in B_TRACE_HEADERS if header not in header_map]
    if missing:
        raise ValueError(f"B 表缺少必要字段：{', '.join(missing)}")

    records: list[BRecord] = []
    for row_number in range(2, worksheet.max_row + 1):
        values = {
            header: worksheet.cell(row=row_number, column=header_map[header]).value
            for header in B_TRACE_HEADERS
        }
        spec = clean_spec(values.get("规格"))
        if not spec:
            continue
        records.append(
            BRecord(
                row_number=row_number,
                date_key=parse_month_day(values.get("系统出入库时间")),
                spec=spec,
                values=values,
            )
        )
    return records


def read_header_map(worksheet: Worksheet) -> dict[str, int]:
    header_map: dict[str, int] = {}
    for col in range(1, worksheet.max_column + 1):
        value = worksheet.cell(row=1, column=col).value
        header = "" if value is None else str(value).strip()
        if header and header not in header_map:
            header_map[header] = col
    return header_map


def ensure_append_headers(worksheet: Worksheet, *, header_row: int = 1) -> int:
    existing_headers = {
        str(worksheet.cell(row=header_row, column=col).value or "").strip(): col
        for col in range(1, worksheet.max_column + 1)
    }

    first_header = APPEND_HEADERS[0]
    if first_header in existing_headers:
        return existing_headers[first_header]

    start_col = find_last_non_empty_header_col(worksheet, header_row=header_row) + 1
    if start_col <= 6:
        start_col = 7

    for offset, header in enumerate(APPEND_HEADERS):
        cell = worksheet.cell(row=header_row, column=start_col + offset, value=header)
        cell.fill = copy.copy(GREEN_FILL)

    return start_col


def trim_trailing_empty_columns(worksheet: Worksheet) -> None:
    last_value_col = 0
    for row in worksheet.iter_rows():
        for cell in row:
            if cell.value not in (None, "") and cell.column > last_value_col:
                last_value_col = cell.column

    if last_value_col and worksheet.max_column > last_value_col:
        worksheet.delete_cols(last_value_col + 1, worksheet.max_column - last_value_col)


def find_last_non_empty_header_col(worksheet: Worksheet, *, header_row: int) -> int:
    for col in range(min(worksheet.max_column, 200), 0, -1):
        if worksheet.cell(row=header_row, column=col).value not in (None, ""):
            return col
    return 0


def infer_c_columns(worksheet: Worksheet) -> CColumns:
    header_aliases = {
        "date_col": ("日期", "送货时间", "系统出入库时间", "入库时间"),
        "spec_col": ("规格", "规格型号", "型号", "产品名称", "商品名称", "商家编码"),
        "qty_col": ("数量", "入库数量"),
        "price_col": ("单价",),
        "amount_col": ("金额",),
    }
    best: tuple[int, dict[str, int]] | None = None

    for row_number in range(1, min(worksheet.max_row, 8) + 1):
        found: dict[str, int] = {}
        for col in range(1, min(worksheet.max_column, 60) + 1):
            header = str(worksheet.cell(row=row_number, column=col).value or "").strip()
            if not header:
                continue
            for field_name, aliases in header_aliases.items():
                if field_name not in found and any(alias == header for alias in aliases):
                    found[field_name] = col
        found = fill_parent_header_columns(worksheet, row_number, found, header_aliases)
        found = fill_date_column_from_data(worksheet, row_number, found)
        has_core_columns = all(
            field_name in found
            for field_name in ("date_col", "spec_col", "qty_col")
        )
        if has_core_columns and len(found) >= 4 and (best is None or len(found) > len(best[1])):
            best = (row_number, found)

    if best is None:
        return CColumns(
            date_col=C_DATE_COL,
            spec_col=C_SPEC_COL,
            qty_col=C_QTY_COL,
            price_col=C_PRICE_COL,
            amount_col=C_AMOUNT_COL,
            first_detail_row=2,
        )

    header_row, found = best
    return CColumns(
        date_col=found.get("date_col", C_DATE_COL),
        spec_col=found.get("spec_col", C_SPEC_COL),
        qty_col=found.get("qty_col", C_QTY_COL),
        price_col=found.get("price_col", C_PRICE_COL),
        amount_col=found.get("amount_col", C_AMOUNT_COL),
        first_detail_row=header_row + 1,
    )


def fill_parent_header_columns(
    worksheet: Worksheet,
    header_row: int,
    found: dict[str, int],
    header_aliases: dict[str, tuple[str, ...]],
) -> dict[str, int]:
    """Fill missing columns from parent header rows in two-line table headers."""

    if all(field_name in found for field_name in header_aliases):
        return found

    merged = dict(found)
    for row_number in range(header_row - 1, 0, -1):
        for col in range(1, min(worksheet.max_column, 60) + 1):
            header = str(worksheet.cell(row=row_number, column=col).value or "").strip()
            if not header:
                continue
            for field_name, aliases in header_aliases.items():
                if field_name not in merged and any(alias == header for alias in aliases):
                    merged[field_name] = col
    return merged


def fill_date_column_from_data(
    worksheet: Worksheet,
    header_row: int,
    found: dict[str, int],
) -> dict[str, int]:
    """Infer the business date column when the header was replaced by a month title."""

    spec_col = found.get("spec_col")
    if spec_col is None:
        return found

    date_col = found.get("date_col")
    if date_col is not None and date_col < spec_col:
        return found

    best_col: int | None = None
    best_score = 0
    for col in range(1, spec_col):
        score = 0
        for row_number in range(header_row + 1, min(worksheet.max_row, header_row + 20) + 1):
            if parse_month_day(worksheet.cell(row=row_number, column=col).value) is not None:
                score += 1
        if score > best_score:
            best_col = col
            best_score = score

    if best_col is None or best_score == 0:
        return found

    merged = dict(found)
    merged["date_col"] = best_col
    return merged


def find_c_detail_rows(worksheet: Worksheet, c_columns: CColumns) -> list[int]:
    detail_rows: list[int] = []
    for row_number in range(c_columns.first_detail_row, worksheet.max_row + 1):
        if is_total_row(worksheet, row_number):
            break
        spec = clean_spec(worksheet.cell(row=row_number, column=c_columns.spec_col).value)
        qty = to_decimal(worksheet.cell(row=row_number, column=c_columns.qty_col).value)
        if spec and qty is not None:
            detail_rows.append(row_number)
    return detail_rows


def find_total_row(worksheet: Worksheet) -> int | None:
    for row_number in range(1, worksheet.max_row + 1):
        if is_total_row(worksheet, row_number):
            return row_number
    return None


def is_total_row(worksheet: Worksheet, row_number: int) -> bool:
    value = worksheet.cell(row=row_number, column=C_DATE_COL).value
    return isinstance(value, str) and "合计" in value


def load_cached_a_values(
    metadata: dict[str, Any],
    c_columns: CColumns,
) -> dict[int, dict[str, Any]]:
    a_file_path = metadata.get("a_file_path")
    if not a_file_path:
        return {}

    path = Path(str(a_file_path))
    if not path.exists():
        return {}

    workbook = load_workbook(path, data_only=True)
    worksheet = workbook.active
    values: dict[int, dict[str, Any]] = {}
    for row_number in range(c_columns.first_detail_row, worksheet.max_row + 1):
        values[row_number] = {
            "price": worksheet.cell(row=row_number, column=c_columns.price_col).value,
            "amount": worksheet.cell(row=row_number, column=c_columns.amount_col).value,
        }
    return values


def materialize_amount_formulas(
    worksheet: Worksheet,
    detail_rows: Iterable[int],
    c_columns: CColumns,
    cached_values: dict[int, dict[str, Any]] | None = None,
) -> None:
    for row_number in detail_rows:
        cached_row = (cached_values or {}).get(row_number, {})
        price_cell = worksheet.cell(row=row_number, column=c_columns.price_col)
        if isinstance(price_cell.value, str) and price_cell.value.startswith("="):
            cached_price = cached_row.get("price")
            if cached_price is not None:
                price_cell.value = cached_price
                price_cell.number_format = "General"

        amount_cell = worksheet.cell(row=row_number, column=c_columns.amount_col)
        if not isinstance(amount_cell.value, str) or not amount_cell.value.startswith("="):
            continue
        cached_amount = cached_row.get("amount")
        if cached_amount is not None:
            amount_cell.value = cached_amount
            amount_cell.number_format = "General"
            continue

        quantity = to_decimal(worksheet.cell(row=row_number, column=c_columns.qty_col).value)
        price = to_decimal(worksheet.cell(row=row_number, column=c_columns.price_col).value)
        if quantity is None or price is None:
            continue
        amount_cell.value = decimal_to_excel_value(quantity * price)
        amount_cell.number_format = "General"


def build_match_plans(
    worksheet: Worksheet,
    detail_rows: list[int],
    mapping_by_a_value: dict[str, list[str]],
    b_records: list[BRecord],
    c_columns: CColumns,
) -> dict[int, MatchPlan]:
    plans: dict[int, MatchPlan] = {}
    used_b_rows: set[int] = set()

    for row_number in detail_rows:
        a_spec = clean_spec(worksheet.cell(row=row_number, column=c_columns.spec_col).value)
        b_specs = mapping_by_a_value.get(a_spec, [])
        quantity = to_decimal(worksheet.cell(row=row_number, column=c_columns.qty_col).value)
        date_key = parse_month_day(worksheet.cell(row=row_number, column=c_columns.date_col).value)

        if not b_specs:
            plans[row_number] = MatchPlan(
                status="未匹配/无规格映射",
                basis=f"A规格未在 metadata.product_mapping.result.mappings 中找到：{a_spec}",
                direction="",
                records=[],
            )
            continue
        if quantity is None:
            plans[row_number] = MatchPlan(
                status="未匹配/数量为空",
                basis="C表数量为空或不是数字",
                direction="",
                records=[],
            )
            continue

        direction = "入库" if quantity > 0 else "出库" if quantity < 0 else ""
        if not direction:
            plans[row_number] = MatchPlan(
                status="零数量需复核",
                basis="C表数量为 0，不能自动判断入库或出库",
                direction="",
                records=[],
            )
            continue

        target = abs(quantity)
        candidates = [
            record
            for record in b_records
            if record.row_number not in used_b_rows
            and record.spec in b_specs
            and record.date_key == date_key
            and record.quantity_for_direction(direction) is not None
        ]
        selected = select_records_by_quantity(candidates, target, direction)

        if selected:
            used_b_rows.update(record.row_number for record in selected)
            plans[row_number] = build_matched_plan(
                worksheet,
                row_number,
                c_columns,
                selected,
                direction,
                status="已匹配" if len(selected) == 1 else "已匹配/多行展开",
                basis=(
                    f"同日期+metadata规格映射+{direction}数量合计一致；"
                    f"A规格={a_spec}；B规格={join_specs(b_specs)}；B行号={join_row_numbers(selected)}"
                ),
            )
            continue

        date_review_candidates = [
            record
            for record in b_records
            if record.row_number not in used_b_rows
            and record.spec in b_specs
            and record.quantity_for_direction(direction) is not None
        ]
        date_review_selected = select_records_by_quantity(
            date_review_candidates,
            target,
            direction,
        )
        if date_review_selected:
            used_b_rows.update(record.row_number for record in date_review_selected)
            plans[row_number] = build_matched_plan(
                worksheet,
                row_number,
                c_columns,
                date_review_selected,
                direction,
                status=(
                    "已匹配/日期需复核"
                    if len(date_review_selected) == 1
                    else "已匹配/多行展开/日期需复核"
                ),
                basis=(
                    f"metadata规格映射+{direction}数量合计一致，但日期不一致；"
                    f"A日期={date_key}；B日期={join_date_keys(date_review_selected)}；"
                    f"A规格={a_spec}；B规格={join_specs(b_specs)}；"
                    f"B行号={join_row_numbers(date_review_selected)}"
                ),
            )
            continue

        if candidates:
            candidate_quantities = [
                str(record.quantity_for_direction(direction)) for record in candidates[:8]
            ]
            basis = (
                f"找到同日期同规格 B 记录，但{direction}数量无法合计为 {target}；"
                f"候选数量={', '.join(candidate_quantities)}"
            )
        else:
            basis = f"未找到同日期同规格且未承接的 B 表{direction}记录；B规格={join_specs(b_specs)}"
        plans[row_number] = MatchPlan(
            status="未匹配",
            basis=basis,
            direction=direction,
            records=[],
        )

    return plans


def build_matched_plan(
    worksheet: Worksheet,
    row_number: int,
    c_columns: CColumns,
    selected: list[BRecord],
    direction: str,
    *,
    status: str,
    basis: str,
) -> MatchPlan:
    amount_diff = calculate_amount_diff(worksheet, row_number, c_columns, selected, direction)
    if amount_diff is not None and amount_diff != 0:
        status = f"{status}/金额需复核"
        basis = f"{basis}；A金额与B成本金额存在差异"
    return MatchPlan(
        status=status,
        basis=basis,
        direction=direction,
        records=selected,
        amount_diff=amount_diff,
    )


def find_unmatched_b_records(
    b_records: list[BRecord],
    plans: dict[int, MatchPlan],
) -> list[BRecord]:
    used_b_rows = {
        record.row_number
        for plan in plans.values()
        for record in plan.records
    }
    return [
        record
        for record in b_records
        if record.row_number not in used_b_rows and infer_record_direction(record) is not None
    ]


def select_records_by_quantity(
    records: list[BRecord],
    target: Decimal,
    direction: str,
) -> list[BRecord]:
    exact = [
        record
        for record in records
        if record.quantity_for_direction(direction) == target
    ]
    if exact:
        return [exact[0]]

    sums: dict[Decimal, list[BRecord]] = {Decimal(0): []}
    for record in records:
        quantity = record.quantity_for_direction(direction)
        if quantity is None or quantity <= 0:
            continue
        snapshot = list(sums.items())
        for current_sum, current_records in snapshot:
            new_sum = current_sum + quantity
            if new_sum in sums:
                continue
            new_records = [*current_records, record]
            if new_sum == target:
                return new_records
            if new_sum < target:
                sums[new_sum] = new_records

    return []


def append_unmatched_b_records(
    worksheet: Worksheet,
    records: list[BRecord],
    c_columns: CColumns,
    *,
    start_col: int,
) -> None:
    if not records:
        return

    total_row = find_total_row(worksheet)
    if total_row is None:
        first_row = worksheet.max_row + 1
    else:
        first_row = total_row
        worksheet.insert_rows(first_row, amount=len(records))

    for index, record in enumerate(records):
        row_number = first_row + index
        direction = infer_record_direction(record)
        if direction is None:
            continue
        quantity = record.quantity_for_direction(direction)
        price_header = "入库成本单价" if direction == "入库" else "出库成本单价"
        amount_header = "入库成本金额" if direction == "入库" else "出库成本金额"
        signed_quantity = quantity if direction == "入库" else -quantity if quantity is not None else None

        worksheet.cell(row=row_number, column=c_columns.date_col, value=record.values.get("系统出入库时间"))
        worksheet.cell(row=row_number, column=c_columns.spec_col, value=record.spec)
        if signed_quantity is not None:
            worksheet.cell(
                row=row_number,
                column=c_columns.qty_col,
                value=decimal_to_excel_value(signed_quantity),
            )
        worksheet.cell(row=row_number, column=c_columns.price_col, value=record.values.get(price_header))
        worksheet.cell(row=row_number, column=c_columns.amount_col, value=record.values.get(amount_header))

        plan = MatchPlan(
            status="B表未匹配追加",
            basis="B表记录未被A表承接，按反向核查追加",
            direction=direction,
            records=[record],
        )
        write_plan_row(worksheet, row_number, plan, record, start_col=start_col)


def infer_record_direction(record: BRecord) -> str | None:
    in_quantity = record.quantity_for_direction("入库")
    if in_quantity is not None and in_quantity != 0:
        return "入库"
    out_quantity = record.quantity_for_direction("出库")
    if out_quantity is not None and out_quantity != 0:
        return "出库"
    return None


def calculate_amount_diff(
    worksheet: Worksheet,
    row_number: int,
    c_columns: CColumns,
    records: list[BRecord],
    direction: str,
) -> Decimal | None:
    a_amount = to_decimal(worksheet.cell(row=row_number, column=c_columns.amount_col).value)
    if a_amount is None:
        return None

    b_amount_header = "入库成本金额" if direction == "入库" else "出库成本金额"
    b_amounts = [to_decimal(record.values.get(b_amount_header)) for record in records]
    if any(amount is None for amount in b_amounts):
        return None

    b_total = sum((amount for amount in b_amounts if amount is not None), Decimal(0))
    return b_total - abs(a_amount)


def apply_match_plans(
    worksheet: Worksheet,
    plans: dict[int, MatchPlan],
    *,
    start_col: int,
) -> None:
    for row_number in sorted(plans, reverse=True):
        plan = plans[row_number]
        records = plan.records
        if len(records) > 1:
            worksheet.insert_rows(row_number + 1, amount=len(records) - 1)
            for inserted_row in range(row_number + 1, row_number + len(records)):
                clear_original_business_cells(worksheet, inserted_row, first_append_col=start_col)

        target_rows = [row_number + index for index in range(max(1, len(records)))]
        records_to_write = records or [None]
        for target_row, record in zip(target_rows, records_to_write, strict=True):
            write_plan_row(
                worksheet,
                target_row,
                plan,
                record,
                start_col=start_col,
            )


def clear_original_business_cells(
    worksheet: Worksheet,
    row_number: int,
    *,
    first_append_col: int,
) -> None:
    for col in range(1, first_append_col):
        worksheet.cell(row=row_number, column=col).value = None


def write_plan_row(
    worksheet: Worksheet,
    row_number: int,
    plan: MatchPlan,
    record: BRecord | None,
    *,
    start_col: int,
) -> None:
    status_fill = (
        GREEN_FILL
        if plan.records and not plan.status.startswith("B表未匹配")
        else YELLOW_FILL
    )
    values: list[Any] = [
        plan.status,
        plan.basis,
        record.row_number if record else None,
        plan.direction,
        decimal_to_excel_value(plan.amount_diff) if plan.amount_diff is not None else None,
    ]

    if record is None:
        values.extend([None] * len(B_TRACE_HEADERS))
    else:
        values.extend(record.values.get(header) for header in B_TRACE_HEADERS)

    for offset, value in enumerate(values):
        cell = worksheet.cell(row=row_number, column=start_col + offset, value=value)
        cell.fill = copy.copy(status_fill)
        if start_col + offset >= start_col + 10:
            cell.number_format = "General"


def rewrite_total_formula(worksheet: Worksheet, c_columns: CColumns) -> None:
    total_row = find_total_row(worksheet)
    if total_row is None or total_row <= 2:
        return

    amount_cell = worksheet.cell(row=total_row, column=c_columns.amount_col)
    amount_column_letter = get_column_letter(c_columns.amount_col)
    amount_cell.value = f"=SUM({amount_column_letter}{c_columns.first_detail_row}:{amount_column_letter}{total_row - 1})"
    amount_cell.number_format = "General"


def reset_sheet_view(worksheet: Worksheet) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.topLeftCell = "A1"
    if worksheet.sheet_view.selection:
        selection = worksheet.sheet_view.selection[0]
        selection.activeCell = "A1"
        selection.sqref = "A1"


def record_match_metadata(
    metadata_path: Path,
    metadata: dict[str, Any],
    *,
    c_path: Path,
    summary: dict[str, int],
) -> None:
    workflow_steps = get_or_create_list(metadata, "workflow_steps")
    workflow_steps.append(
        {
            "step": "02_match_c_to_b",
            "status": "DONE",
            "c_file_path": str(c_path),
            "summary": summary,
            "finished_at": datetime.now(UTC).isoformat(timespec="seconds"),
        }
    )
    metadata_path.write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")


def has_done_step(metadata: dict[str, Any], step_name: str) -> bool:
    steps = metadata.get("workflow_steps")
    if not isinstance(steps, list):
        return False
    return any(
        isinstance(step, dict)
        and step.get("step") == step_name
        and step.get("status") == "DONE"
        for step in steps
    )


def summarize_plans(plans: dict[int, MatchPlan]) -> dict[str, int]:
    matched_source_rows = 0
    expanded_b_rows = 0
    unmatched_source_rows = 0
    review_source_rows = 0

    for plan in plans.values():
        if plan.records:
            matched_source_rows += 1
            expanded_b_rows += len(plan.records)
            continue
        if "复核" in plan.status:
            review_source_rows += 1
        else:
            unmatched_source_rows += 1

    return {
        "source_rows": len(plans),
        "matched_source_rows": matched_source_rows,
        "expanded_b_rows": expanded_b_rows,
        "unmatched_source_rows": unmatched_source_rows,
        "review_source_rows": review_source_rows,
    }


def clean_spec(value: object) -> str:
    return excel_tool.clean_product_text(value)


def parse_month_day(value: object) -> tuple[int, int] | None:
    if isinstance(value, datetime):
        return value.month, value.day
    if isinstance(value, date):
        return value.month, value.day
    if isinstance(value, int | float) and not isinstance(value, bool):
        return parse_numeric_month_day(float(value))
    if isinstance(value, str):
        text = value.strip()
        match = re.search(r"(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", text)
        if match:
            return int(match.group(2)), int(match.group(3))
        match = re.search(r"(\d{1,2})\s*月\s*(\d{1,2})\s*日?", text)
        if match:
            return int(match.group(1)), int(match.group(2))
    return None


def to_decimal(value: object) -> Decimal | None:
    if value is None or value == "":
        return None
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError):
        return None


def decimal_to_excel_value(value: Decimal) -> int | float:
    if value == value.to_integral_value():
        return int(value)
    return float(value)


def join_row_numbers(records: Iterable[BRecord]) -> str:
    return ",".join(str(record.row_number) for record in records)


def join_specs(specs: Iterable[str]) -> str:
    return "、".join(specs)


def join_date_keys(records: Iterable[BRecord]) -> str:
    return ",".join(str(record.date_key) for record in records)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="按 metadata 映射把 B 表记录承接到 C 表。")
    parser.add_argument("--metadata-json", type=Path, help="上传任务 metadata.json 路径")
    parser.add_argument("--task-dir", type=Path, help="上传任务目录，目录内应包含 metadata.json")
    parser.add_argument("--force", action="store_true", help="即使 metadata 已记录第二步完成也强制执行")
    return parser


def main() -> None:
    args = build_parser().parse_args()
    if not args.metadata_json and not args.task_dir:
        raise SystemExit("请使用 --metadata-json 或 --task-dir")
    metadata_path = args.metadata_json or args.task_dir / "metadata.json"
    output_path = match_c_to_b_from_metadata(metadata_path, force=args.force)
    print(f"C 表已完成 B 表匹配：{output_path}")


if __name__ == "__main__":
    main()
