"""Copy uploaded A workbook into the initial C workbook draft.

This module is used after the upload flow has saved A/B workbooks and written
``metadata.json``. The numbered ``01_copy_a_to_c.py`` script is only a command
entry for the same logic.
"""

from __future__ import annotations

import argparse
import copy
import json
import re
import shutil
from collections.abc import Iterable
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.datetime import from_excel
from openpyxl.worksheet.worksheet import Worksheet

DATE_TEXT_PATTERNS = (
    re.compile(r"^\s*(?:\d{4})[./-](\d{1,2})[./-](\d{1,2})\s*$"),
    re.compile(r"^\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日?\s*$"),
)


def copy_a_to_c(
    a_path: Path,
    c_path: Path,
    *,
    sheet_name: str | None = None,
    date_column: str = "A",
    header_row: int = 1,
    month_title: str | None = None,
    clear_fill: bool = True,
    overwrite: bool = False,
) -> Path:
    """Create a C workbook draft from A workbook."""

    a_path = a_path.resolve()
    c_path = c_path.resolve()
    validate_paths(a_path, c_path, overwrite=overwrite)

    c_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(a_path, c_path)

    workbook = load_workbook(c_path)
    worksheets = resolve_worksheets(workbook.worksheets, sheet_name)
    date_column_index = column_index_from_string(date_column)

    for worksheet in worksheets:
        unmerge_and_fill_dates(
            worksheet,
            date_column_index=date_column_index,
            workbook_epoch=workbook.epoch,
        )
        normalize_date_column(
            worksheet,
            date_column_index=date_column_index,
            workbook_epoch=workbook.epoch,
            start_row=header_row + 1,
        )
        if month_title:
            worksheet.cell(row=header_row, column=date_column_index, value=month_title)
        if clear_fill:
            clear_sheet_fills(worksheet)
        reset_sheet_view(worksheet, header_row=header_row)

    workbook.save(c_path)
    return c_path


def copy_a_to_c_from_metadata(
    metadata_path: Path,
    *,
    output_name: str = "C.xlsx",
    overwrite: bool = True,
    update_metadata: bool = True,
) -> Path:
    """Create the C draft from a task ``metadata.json`` generated after upload."""

    metadata_path = metadata_path.resolve()
    metadata = load_metadata(metadata_path)
    task_dir = metadata_path.parent
    a_path = Path(require_string(metadata, "a_file_path"))
    c_path = Path(str(metadata.get("c_file_path") or task_dir / output_name))

    product_mapping = get_dict(metadata, "product_mapping")
    structure = get_dict(product_mapping, "structure")
    a_sheet_structure = get_dict(structure, "a_sheet")
    a_preview = get_dict(metadata, "a_preview")

    header_row = parse_positive_int(
        a_sheet_structure.get("header_row_guess"),
        default=parse_positive_int(get_dict(metadata, "llm_result").get("header_row_guess"), default=1),
    )
    sheet_name = parse_optional_string(a_preview.get("sheet_name"))
    date_column = infer_date_column_from_preview(a_preview, header_row=header_row)
    month_title = month_title_from_metadata(metadata)

    output_path = copy_a_to_c(
        a_path,
        c_path,
        sheet_name=sheet_name,
        date_column=date_column,
        header_row=header_row,
        month_title=month_title,
        overwrite=overwrite,
    )

    if update_metadata:
        record_c_draft_metadata(
            metadata_path,
            metadata,
            c_path=output_path,
            sheet_name=sheet_name,
            date_column=date_column,
            header_row=header_row,
            month_title=month_title,
        )

    return output_path


def load_metadata(metadata_path: Path) -> dict[str, Any]:
    if metadata_path.suffix.lower() != ".json":
        raise ValueError(f"任务元数据必须是 JSON 文件：{metadata_path}")
    if not metadata_path.exists():
        raise FileNotFoundError(f"任务元数据不存在：{metadata_path}")

    payload = json.loads(metadata_path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise TypeError(f"任务元数据格式无效：{metadata_path}")
    return payload


def record_c_draft_metadata(
    metadata_path: Path,
    metadata: dict[str, Any],
    *,
    c_path: Path,
    sheet_name: str | None,
    date_column: str,
    header_row: int,
    month_title: str | None,
) -> None:
    metadata["c_file_path"] = str(c_path)
    workflow_steps = get_or_create_list(metadata, "workflow_steps")
    workflow_steps.append(
        {
            "step": "01_copy_a_to_c",
            "status": "DONE",
            "c_file_path": str(c_path),
            "sheet_name": sheet_name,
            "date_column": date_column,
            "header_row": header_row,
            "month_title": month_title,
            "finished_at": datetime.now(UTC).isoformat(timespec="seconds"),
        }
    )
    metadata_path.write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")


def validate_paths(a_path: Path, c_path: Path, *, overwrite: bool) -> None:
    if a_path.suffix.lower() != ".xlsx":
        raise ValueError(f"A 表必须是 .xlsx 文件：{a_path}")
    if c_path.suffix.lower() != ".xlsx":
        raise ValueError(f"C 表必须是 .xlsx 文件：{c_path}")
    if not a_path.exists():
        raise FileNotFoundError(f"A 表不存在：{a_path}")
    if a_path == c_path:
        raise ValueError("C 表输出路径不能和 A 表原始路径相同")
    if c_path.exists() and not overwrite:
        raise FileExistsError(f"C 表已存在，若要覆盖请加 --overwrite：{c_path}")


def resolve_worksheets(worksheets: Iterable[Worksheet], sheet_name: str | None) -> list[Worksheet]:
    worksheet_list = list(worksheets)
    if sheet_name is None:
        return worksheet_list

    for worksheet in worksheet_list:
        if worksheet.title == sheet_name:
            return [worksheet]

    available = ", ".join(worksheet.title for worksheet in worksheet_list)
    raise ValueError(f"工作表不存在：{sheet_name}；可用工作表：{available}")


def unmerge_and_fill_dates(
    worksheet: Worksheet,
    *,
    date_column_index: int,
    workbook_epoch: datetime,
) -> None:
    """Unmerge all cells and fill vertical date merges in the date column."""

    merged_ranges = list(worksheet.merged_cells.ranges)
    for merged_range in merged_ranges:
        min_col, min_row, max_col, max_row = merged_range.bounds
        top_left = worksheet.cell(row=min_row, column=min_col)
        top_left_value = top_left.value
        top_left_style = copy_cell_style(top_left)
        worksheet.unmerge_cells(str(merged_range))

        is_date_column_merge = (
            min_col == date_column_index and max_col == date_column_index and max_row > min_row
        )
        if not is_date_column_merge:
            continue

        display_value = format_date_display(top_left_value, workbook_epoch)
        for row_number in range(min_row, max_row + 1):
            cell = worksheet.cell(row=row_number, column=date_column_index)
            cell.value = display_value
            apply_cell_style(cell, top_left_style)


def normalize_date_column(
    worksheet: Worksheet,
    *,
    date_column_index: int,
    workbook_epoch: datetime,
    start_row: int,
) -> None:
    for row_number in range(max(start_row, 1), worksheet.max_row + 1):
        cell = worksheet.cell(row=row_number, column=date_column_index)
        if cell.value is None:
            continue
        cell.value = format_date_display(cell.value, workbook_epoch)


def format_date_display(value: object, workbook_epoch: datetime) -> object:
    if isinstance(value, datetime):
        return f"{value.month}月{value.day}日"
    if isinstance(value, date):
        return f"{value.month}月{value.day}日"
    if isinstance(value, int | float) and not isinstance(value, bool) and 1 <= value <= 60000:
        try:
            converted = from_excel(value, workbook_epoch)
        except (TypeError, ValueError):
            return value
        return f"{converted.month}月{converted.day}日"
    if isinstance(value, str):
        stripped = value.strip()
        for pattern in DATE_TEXT_PATTERNS:
            match = pattern.match(stripped)
            if match:
                month, day = match.groups()
                return f"{int(month)}月{int(day)}日"
        return stripped
    return value


def infer_date_column_from_preview(a_preview: dict[str, Any], *, header_row: int) -> str:
    rows = a_preview.get("rows")
    if not isinstance(rows, list):
        return "A"

    header_cells: list[dict[str, Any]] = []
    for row in rows:
        if not isinstance(row, dict):
            continue
        if parse_positive_int(row.get("row_number"), default=0) == header_row:
            cells = row.get("cells")
            if isinstance(cells, list):
                header_cells = [cell for cell in cells if isinstance(cell, dict)]
            break

    for cell in header_cells:
        value = parse_optional_string(cell.get("value"))
        if value and ("日期" in value or "时间" in value):
            return get_column_letter(parse_positive_int(cell.get("column"), default=1))

    return "A"


def month_title_from_metadata(metadata: dict[str, Any]) -> str | None:
    reconcile_month = parse_optional_string(metadata.get("reconcile_month"))
    if not reconcile_month:
        return None

    match = re.match(r"^\d{4}-(\d{1,2})$", reconcile_month)
    if not match:
        return None

    month_names = {
        1: "一月",
        2: "二月",
        3: "三月",
        4: "四月",
        5: "五月",
        6: "六月",
        7: "七月",
        8: "八月",
        9: "九月",
        10: "十月",
        11: "十一月",
        12: "十二月",
    }
    return month_names.get(int(match.group(1)))


def copy_cell_style(cell: Any) -> dict[str, object]:
    return {
        "font": copy.copy(cell.font),
        "fill": copy.copy(cell.fill),
        "border": copy.copy(cell.border),
        "alignment": copy.copy(cell.alignment),
        "number_format": cell.number_format,
        "protection": copy.copy(cell.protection),
    }


def apply_cell_style(cell: Any, style: dict[str, object]) -> None:
    cell.font = style["font"]
    cell.fill = style["fill"]
    cell.border = style["border"]
    cell.alignment = style["alignment"]
    cell.number_format = style["number_format"]
    cell.protection = style["protection"]


def clear_sheet_fills(worksheet: Worksheet) -> None:
    empty_fill = PatternFill(fill_type=None)
    for row in worksheet.iter_rows():
        for cell in row:
            cell.fill = copy.copy(empty_fill)


def reset_sheet_view(worksheet: Worksheet, *, header_row: int) -> None:
    worksheet.freeze_panes = f"A{header_row + 1}"
    worksheet.sheet_view.topLeftCell = "A1"
    if worksheet.sheet_view.selection:
        selection = worksheet.sheet_view.selection[0]
        selection.activeCell = "A1"
        selection.sqref = "A1"


def get_dict(payload: dict[str, Any], key: str) -> dict[str, Any]:
    value = payload.get(key)
    if isinstance(value, dict):
        return value
    return {}


def get_or_create_list(payload: dict[str, Any], key: str) -> list[Any]:
    value = payload.get(key)
    if isinstance(value, list):
        return value

    value = []
    payload[key] = value
    return value


def require_string(payload: dict[str, Any], key: str) -> str:
    value = parse_optional_string(payload.get(key))
    if value is None:
        raise ValueError(f"任务元数据缺少字段：{key}")
    return value


def parse_optional_string(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def parse_positive_int(value: object, *, default: int) -> int:
    try:
        number = int(value)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return default
    return number if number > 0 else default


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="上传任务 JSON 生成 C 表底稿。")
    parser.add_argument("a_path", nargs="?", type=Path, help="A 表 .xlsx 路径")
    parser.add_argument("c_path", nargs="?", type=Path, help="输出 C 表 .xlsx 路径")
    parser.add_argument("--metadata-json", type=Path, help="上传流程生成的 metadata.json 路径")
    parser.add_argument("--task-dir", type=Path, help="上传任务目录，目录内应包含 metadata.json")
    parser.add_argument("--sheet-name", help="只处理指定工作表；不传则使用 JSON 中的 A 表工作表")
    parser.add_argument("--date-column", default="A", help="日期列，默认 A")
    parser.add_argument("--header-row", type=int, default=1, help="表头行号，默认 1")
    parser.add_argument("--month-title", help="写入表头日期列的月份标题，例如 七月")
    parser.add_argument("--keep-fill", action="store_true", help="保留 A 表原底色")
    parser.add_argument("--overwrite", action="store_true", help="允许覆盖已存在的 C 表")
    return parser


def main() -> None:
    args = build_parser().parse_args()
    if args.task_dir or args.metadata_json:
        metadata_path = args.metadata_json or args.task_dir / "metadata.json"
        output_path = copy_a_to_c_from_metadata(metadata_path, overwrite=True)
    else:
        if args.a_path is None or args.c_path is None:
            raise SystemExit("请传入 A/C 路径，或使用 --metadata-json / --task-dir")
        output_path = copy_a_to_c(
            args.a_path,
            args.c_path,
            sheet_name=args.sheet_name,
            date_column=args.date_column,
            header_row=args.header_row,
            month_title=args.month_title,
            clear_fill=not args.keep_fill,
            overwrite=args.overwrite,
        )
    print(f"C 表底稿已生成：{output_path}")


if __name__ == "__main__":
    main()
