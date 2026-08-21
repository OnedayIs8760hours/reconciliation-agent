"""工作表版式探测。"""

from __future__ import annotations

from dataclasses import dataclass
from typing import Any

from openpyxl.worksheet.worksheet import Worksheet

from agent.rules.rule_schema import RuleConfig


@dataclass(slots=True)
class SheetProfile:
    sheet_name: str
    header_row: int
    columns: dict[str, int]
    max_row: int
    max_column: int


def _norm_header(value: Any) -> str:
    return str(value or "").strip().replace(" ", "").replace("\n", "")


class WorkbookProfiler:
    def __init__(self, rule_config: RuleConfig) -> None:
        self.rule_config = rule_config

    def profile(self, worksheet: Worksheet) -> SheetProfile:
        best_row = 1
        best_columns: dict[str, int] = {}
        for row in range(1, min(20, worksheet.max_row) + 1):
            row_values = [_norm_header(worksheet.cell(row=row, column=column).value) for column in range(1, worksheet.max_column + 1)]
            columns: dict[str, int] = {}
            for field, aliases in self.rule_config.header_aliases.items():
                for idx, header in enumerate(row_values, start=1):
                    if header and any(alias.replace(" ", "") in header for alias in aliases):
                        columns.setdefault(field, idx)
                        break
            if len(columns) > len(best_columns):
                best_row = row
                best_columns = columns
        return SheetProfile(
            sheet_name=worksheet.title,
            header_row=best_row,
            columns=best_columns,
            max_row=worksheet.max_row,
            max_column=worksheet.max_column,
        )
