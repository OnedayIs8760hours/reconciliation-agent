"""标记未被 C 表验收覆盖的 B 表记录。"""

from pathlib import Path

from openpyxl.styles import PatternFill

from agent.core.workbook_loader import WorkbookLoader
from agent.domain.records import BRecord, MissingBRecord
from agent.domain.statuses import MatchDirection, MissingType


class MissingMarker:
    """创建带有明确缺失注释的 B 表副本。"""

    status_header = "C表反向核查状态"
    note_header = "C表反向核查说明"

    def __init__(self, loader: WorkbookLoader) -> None:
        self.loader = loader

    def infer_missing_records(self, uncovered_b_records: list[BRecord]) -> list[MissingBRecord]:
        """为未被 C 表接受的 B 表行构建缺失注释。"""
        missing: list[MissingBRecord] = []
        for record in uncovered_b_records:
            if record.inbound_qty is not None:
                direction = MatchDirection.INBOUND
                missing_type = MissingType.INBOUND_MISSING
            elif record.outbound_qty is not None:
                direction = MatchDirection.OUTBOUND
                missing_type = MissingType.OUTBOUND_MISSING
            else:
                direction = MatchDirection.ZERO_OR_SPECIAL
                missing_type = MissingType.ZERO_AMOUNT_MISSING
            missing.append(
                MissingBRecord(
                    b_record=record,
                    direction=direction,
                    missing_type=missing_type,
                    basis="B表本月记录未在C表承接，按强制规则标注缺失",
                    note=f"{direction.value}记录未进入C表，需补充或人工排除",
                )
            )
        return missing

    def mark(self, b_table_path: Path, missing_records: list[MissingBRecord], output_path: Path) -> Path:
        """将缺失状态和备注写入 B 表副本。"""
        workbook = self.loader.load(b_table_path)
        sheet = self.loader.active_sheet(workbook)
        status_col = sheet.max_column + 1
        note_col = status_col + 1
        sheet.cell(1, status_col, self.status_header)
        sheet.cell(1, note_col, self.note_header)
        fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
        for missing in missing_records:
            row = missing.b_record.row_number
            sheet.cell(row, status_col, missing.missing_type.value)
            sheet.cell(row, note_col, missing.note)
            for col_idx in range(1, note_col + 1):
                sheet.cell(row, col_idx).fill = fill
        return self.loader.save(workbook, output_path)
