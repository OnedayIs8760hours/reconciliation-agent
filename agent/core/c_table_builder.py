"""基于 A 表生成 C 表草稿。"""

from copy import copy
from pathlib import Path

from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from agent.domain.schemas import CAppendColumns
from agent.core.workbook_loader import WorkbookLoader


class CTableBuilder:
    """通过完整复制 A 表并追加追踪列来构建初始 C 表。"""

    def __init__(self, loader: WorkbookLoader, append_columns: CAppendColumns | None = None) -> None:
        self.loader = loader
        self.append_columns = append_columns or CAppendColumns()

    def build_from_a_table(self, a_table_path: Path, output_path: Path) -> Path:
        """从 A 表创建 C 表，保留原始业务列和样式。"""
        workbook = self.loader.load(a_table_path)
        sheet = self.loader.active_sheet(workbook)
        self._unmerge_and_fill(sheet)
        self._append_trace_headers(sheet)
        self._normalize_date_display(sheet)
        return self.loader.save(workbook, output_path)

    def _unmerge_and_fill(self, sheet: Worksheet) -> None:
        """取消合并单元格，并将合并值填入原先合并区域中的每个单元格。"""
        ranges = list(sheet.merged_cells.ranges)
        for merged_range in ranges:
            value = sheet.cell(merged_range.min_row, merged_range.min_col).value
            sheet.unmerge_cells(str(merged_range))
            for row in sheet.iter_rows(
                min_row=merged_range.min_row,
                max_row=merged_range.max_row,
                min_col=merged_range.min_col,
                max_col=merged_range.max_col,
            ):
                for cell in row:
                    cell.value = value

    def _append_trace_headers(self, sheet: Worksheet) -> None:
        headers = [cell.value for cell in sheet[1]]
        next_column = sheet.max_column + 1
        for header in self.append_columns.__dict__.values():
            if header not in headers:
                cell = sheet.cell(1, next_column, header)
                if sheet.max_column >= 1:
                    source = sheet.cell(1, max(1, next_column - 1))
                    cell.font = copy(source.font)
                    cell.fill = copy(source.fill)
                    cell.border = copy(source.border)
                    cell.alignment = copy(source.alignment)
                next_column += 1

    def _normalize_date_display(self, sheet: Worksheet) -> None:
        """复制 A 表后保持 Excel 日期显示稳定。"""
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if cell.is_date:
                    cell.number_format = "yyyy/m/d"

    def clone_inserted_row_style(self, sheet: Worksheet, source_row: int, target_row: int) -> None:
        """将一行的样式复制到插入的扩展行。"""
        for col_idx in range(1, sheet.max_column + 1):
            source = sheet.cell(source_row, col_idx)
            target = sheet.cell(target_row, col_idx)
            target.font = copy(source.font)
            target.fill = copy(source.fill)
            target.border = copy(source.border)
            target.alignment = copy(source.alignment)
            target.number_format = source.number_format
