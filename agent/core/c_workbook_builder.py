"""C 表工作簿生成。"""

from __future__ import annotations

from copy import copy
from pathlib import Path

from openpyxl import load_workbook


class CWorkbookBuilder:
    def build_from_a(self, a_path: Path, c_path: Path):
        workbook = load_workbook(a_path)
        worksheet = workbook.active
        for merged_range in list(worksheet.merged_cells.ranges):
            top_left = worksheet.cell(row=merged_range.min_row, column=merged_range.min_col)
            value = top_left.value
            style = copy(top_left._style)
            worksheet.unmerge_cells(str(merged_range))
            for row in range(merged_range.min_row, merged_range.max_row + 1):
                for column in range(merged_range.min_col, merged_range.max_col + 1):
                    cell = worksheet.cell(row=row, column=column)
                    cell.value = value
                    cell._style = copy(style)
        workbook.save(c_path)
        return workbook, worksheet
