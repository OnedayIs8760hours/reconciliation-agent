from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from backend.domain.reconciliation_record import BTableRecord
from backend.domain.reconciliation_schema import BSheetFieldSchema
from backend.services.reverse_verify_service import ReverseVerifySummary
from tools import excel_tool


B_MISSING_HEADERS = [
    "C表缺失标注",
    "缺失类型",
    "人工表行号",
    "缺失匹配依据",
    "标注说明",
]


@dataclass(frozen=True)
class BMissingMarkerResult:
    """B 表缺失标注结果。"""

    file_path: str
    marked_count: int
    marker_start_column: int
    marker_headers: list[str]
    sheet_name: str = ""

    def to_dict(self) -> dict[str, object]:
        """把 B 表缺失标注结果转成 JSON 字典。"""

        return {
            "file_path": self.file_path,
            "marked_count": self.marked_count,
            "marker_start_column": self.marker_start_column,
            "marker_headers": list(self.marker_headers),
            "sheet_name": self.sheet_name,
        }


class BMissingMarkerService:
    """复制 B 表并在本月未承接记录右侧写缺失标注。"""

    def mark_missing_records(
        self,
        b_file_path: Path,
        marked_file_path: Path,
        b_schema: BSheetFieldSchema,
        reverse_summary: ReverseVerifySummary,
    ) -> BMissingMarkerResult:
        """把反向核查发现的缺失 B 记录标注到 B_marked.xlsx。"""

        # copy_workbook(...)：先完整复制 B 表，避免修改用户上传的 B 表原件。
        excel_tool.copy_workbook(b_file_path, marked_file_path)
        # load_workbook(...)：打开复制后的 B 标注表。
        workbook = load_workbook(marked_file_path)
        worksheet = excel_tool.get_sheet(workbook, b_schema.sheet_name or None)

        marker_start_column = worksheet.max_column + 1
        self.write_marker_headers(worksheet, b_schema.header_row, marker_start_column)
        self.write_missing_rows(worksheet, marker_start_column, reverse_summary.missing_records)

        excel_tool.reset_view(worksheet, header_row=b_schema.header_row)
        excel_tool.auto_fit_columns(worksheet)
        # save(...)：保存 B 表缺失标注结果。
        workbook.save(marked_file_path)

        return BMissingMarkerResult(
            file_path=str(marked_file_path),
            marked_count=len(reverse_summary.missing_records),
            marker_start_column=marker_start_column,
            marker_headers=list(B_MISSING_HEADERS),
            sheet_name=worksheet.title,
        )

    def write_marker_headers(self, worksheet: Worksheet, header_row: int, start_column: int) -> None:
        """写入 B 表缺失标注表头。"""

        for index, header in enumerate(B_MISSING_HEADERS):
            # cell(...)：按行列号把标注表头写到 B 表右侧。
            worksheet.cell(row=header_row, column=start_column + index, value=header)
        excel_tool.style_header(
            worksheet,
            row=header_row,
            start_col=start_column,
            end_col=start_column + len(B_MISSING_HEADERS) - 1,
        )

    def write_missing_rows(self, worksheet: Worksheet, start_column: int, missing_records: list[BTableRecord]) -> None:
        """在每条缺失 B 表记录所在行写入标注信息。"""

        yellow_fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
        for record in missing_records:
            values = [
                "C表缺失",
                "B表本月记录未在C表承接",
                "",
                record.trace_key(),
                "该 B 表本月系统记录没有匹配到 A/C 表明细，请人工确认是否漏入人工表",
            ]
            for index, value in enumerate(values):
                # cell(...)：定位缺失记录原始行右侧标注列。
                cell = worksheet.cell(row=record.source_row, column=start_column + index, value=value)
                cell.fill = yellow_fill
