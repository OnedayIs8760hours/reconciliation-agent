from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from backend.domain.match_result import MatchSummary
from backend.domain.verify_report import VerifyItem, VerifyReport
from backend.services.b_missing_marker_service import BMissingMarkerResult
from backend.services.c_table_base_service import CTableBaseResult
from backend.services.reverse_verify_service import ReverseVerifySummary
from tools import excel_tool


class ReconciliationVerifyService:
    """执行交付前强制验收，确认 C 表和 B 表反向核查结果可靠。"""

    def verify(
        self,
        a_file_path: Path,
        c_file_path: Path,
        b_marked_file_path: Path,
        c_table_result: CTableBaseResult,
        match_summary: MatchSummary,
        reverse_summary: ReverseVerifySummary,
        marker_result: BMissingMarkerResult,
        quantity_column: int,
        amount_column: int,
        report_file_path: Path,
    ) -> VerifyReport:
        """执行全部强制验收项，并写出 verify_report.json。"""

        items: list[VerifyItem] = []

        a_quantity_total = self.sum_column(
            a_file_path,
            c_table_result.sheet_name,
            c_table_result.detail_start_row,
            c_table_result.detail_end_row,
            quantity_column,
        )
        c_quantity_total = self.sum_column(
            c_file_path,
            c_table_result.sheet_name,
            c_table_result.detail_start_row,
            c_table_result.detail_end_row + c_table_result_to_inserted_rows(match_summary),
            quantity_column,
        )
        quantity_diff = round(c_quantity_total - a_quantity_total, 2)
        items.append(
            VerifyItem(
                name="A/C 原始业务数量合计一致",
                passed=abs(quantity_diff) <= 0.01,
                value=quantity_diff,
                message="C 表插行后原始业务数量合计必须和 A 表一致",
            )
        )

        a_amount_total = self.sum_column(
            a_file_path,
            c_table_result.sheet_name,
            c_table_result.detail_start_row,
            c_table_result.detail_end_row,
            amount_column,
        )
        c_amount_total = self.sum_column(
            c_file_path,
            c_table_result.sheet_name,
            c_table_result.detail_start_row,
            c_table_result.detail_end_row + c_table_result_to_inserted_rows(match_summary),
            amount_column,
        )
        amount_diff = round(c_amount_total - a_amount_total, 2)
        items.append(
            VerifyItem(
                name="A/C 原始业务金额合计一致",
                passed=abs(amount_diff) <= 0.01,
                value=amount_diff,
                message="C 表插行后原始业务金额合计必须和 A 表一致",
            )
        )

        coverage_value = {
            "monthly_records": reverse_summary.monthly_records,
            "accepted_by_c": reverse_summary.accepted_by_c,
            "marked_missing": reverse_summary.marked_missing,
            "manual_excluded": reverse_summary.manually_excluded,
        }
        coverage_total = reverse_summary.accepted_by_c + reverse_summary.marked_missing + reverse_summary.manually_excluded
        items.append(
            VerifyItem(
                name="B 表本月记录全部解释",
                passed=coverage_total == reverse_summary.monthly_records and reverse_summary.unexplained == 0,
                value=coverage_value,
                message="B 表本月记录必须等于 C 表承接 + B 表标注缺失 + 人工排除",
            )
        )

        items.append(
            VerifyItem(
                name="B 表缺失标注数量一致",
                passed=marker_result.marked_count == reverse_summary.marked_missing,
                value={"marker_count": marker_result.marked_count, "reverse_missing": reverse_summary.marked_missing},
                message="反向核查缺失记录必须全部写入 B_marked.xlsx",
            )
        )

        formula_errors = self.scan_formula_errors(
            c_file_path,
            b_marked_file_path,
            c_table_result.sheet_name,
            marker_result.sheet_name,
        )
        items.append(
            VerifyItem(
                name="Excel 公式错误数量为 0",
                passed=len(formula_errors) == 0,
                value=len(formula_errors),
                message="交付文件中不能出现 #REF!、#DIV/0!、#VALUE!、#NAME?、#N/A",
            )
        )

        exceptions = self.build_exceptions(match_summary, reverse_summary)
        passed = True
        for item in items:
            if not item.passed:
                passed = False

        report = VerifyReport(
            passed=passed,
            items=items,
            formula_errors=formula_errors,
            exceptions=exceptions,
        )

        # json.dumps(...)：把验收报告转成 JSON 文本；default=str 用来兼容日期时间值。
        report_text = json.dumps(report.to_dict(), ensure_ascii=False, indent=2, default=str)
        # write_text(...)：把验收报告写入任务目录，方便前端下载和人工排查。
        report_file_path.write_text(report_text, encoding="utf-8")
        return report

    def sum_column(self, file_path: Path, sheet_name: str, start_row: int, end_row: int, column_index: int) -> float:
        """统计指定 Excel 文件中某一列的数字合计。"""

        if column_index < 1:
            return 0.0

        # load_workbook(..., data_only=True)：读取公式缓存值，避免把公式文本当成金额。
        workbook = load_workbook(file_path, data_only=True)
        worksheet = excel_tool.get_sheet(workbook, sheet_name or None)
        total = 0.0
        for row_index in range(start_row, end_row + 1):
            # cell(...)：按行列号读取数量或金额单元格。
            value = worksheet.cell(row=row_index, column=column_index).value
            total = total + to_float(value)
        return total

    def scan_formula_errors(
        self,
        c_file_path: Path,
        b_marked_file_path: Path,
        c_sheet_name: str,
        b_sheet_name: str,
    ) -> list[str]:
        """扫描 C 表和 B 标注表中的公式错误。"""

        errors: list[str] = []
        # load_workbook(..., data_only=False)：读取公式文本和错误文本，方便发现 Excel 错误值。
        c_workbook = load_workbook(c_file_path, data_only=False)
        c_worksheet = excel_tool.get_sheet(c_workbook, c_sheet_name or None)
        for error in excel_tool.scan_formula_errors(c_worksheet):
            errors.append(f"C表 {error}")

        # load_workbook(..., data_only=False)：读取 B 标注表的真实工作表，不能默认 active。
        b_workbook = load_workbook(b_marked_file_path, data_only=False)
        b_worksheet = excel_tool.get_sheet(b_workbook, b_sheet_name or None)
        for error in excel_tool.scan_formula_errors(b_worksheet):
            errors.append(f"B标注表 {error}")

        return errors

    def build_exceptions(self, match_summary: MatchSummary, reverse_summary: ReverseVerifySummary) -> list[dict[str, object]]:
        """把未匹配 A 行和缺失 B 行整理成前端异常列表。"""

        exceptions: list[dict[str, object]] = []
        for result in match_summary.results:
            if result.status != "未匹配" and result.status != "异常":
                continue
            exceptions.append(
                {
                    "id": f"A-{result.a_row}",
                    "source": "A表",
                    "row": result.a_row,
                    "systemTime": "",
                    "documentNo": "",
                    "sku": "",
                    "productName": "",
                    "spec": "",
                    "quantity": 0,
                    "unitCost": 0,
                    "amount": 0,
                    "type": result.status,
                    "reason": result.reason or result.basis,
                }
            )

        for record in reverse_summary.missing_records:
            exceptions.append(
                {
                    "id": f"B-{record.source_row}-{record.direction}",
                    "source": "B表",
                    "row": record.source_row,
                    "systemTime": str(record.system_time) if record.system_time is not None else "",
                    "documentNo": record.document_no,
                    "sku": record.sku,
                    "productName": record.product_name,
                    "spec": record.spec,
                    "quantity": record.quantity,
                    "unitCost": record.unit_price,
                    "amount": record.amount,
                    "type": "B表本月记录未承接",
                    "reason": "该 B 表记录没有被 C 表承接，已在 B_marked.xlsx 标注缺失",
                }
            )

        return exceptions


def c_table_result_to_inserted_rows(match_summary: MatchSummary) -> int:
    """读取 C 表匹配过程中新增的行数。"""

    return match_summary.inserted_rows


def to_float(value: object) -> float:
    """把 Excel 单元格值安全转换成数字。"""

    if value is None or value == "":
        return 0.0
    try:
        # float(...)：把数字或数字字符串转换成小数。
        return float(value)
    except (TypeError, ValueError):
        return 0.0
