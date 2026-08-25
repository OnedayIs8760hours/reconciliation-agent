from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from backend.domain.reconciliation_schema import SheetFieldSchema
from tools import excel_tool


C_TRACE_HEADERS = [
    "匹配状态",
    "匹配依据",
    "B表行号",
    "匹配方向",
    "仓库",
    "系统出入库时间",
    "单据编号",
    "货品编号",
    "货品名称",
    "规格",
    "入库数量",
    "入库成本单价",
    "入库成本金额",
    "出库数量",
    "出库成本单价",
    "出库成本金额",
    "复核类型",
    "处理建议",
]


@dataclass(frozen=True)
class CTableBaseResult:
    """C 表底稿生成结果。"""

    file_path: str
    sheet_name: str
    original_max_row: int
    original_max_column: int
    header_row: int
    detail_start_row: int
    detail_end_row: int
    trace_start_column: int
    trace_headers: list[str]

    def to_dict(self) -> dict[str, object]:
        """把 C 表底稿结果转成 JSON 字典。"""

        return {
            "file_path": self.file_path,
            "sheet_name": self.sheet_name,
            "original_max_row": self.original_max_row,
            "original_max_column": self.original_max_column,
            "header_row": self.header_row,
            "detail_start_row": self.detail_start_row,
            "detail_end_row": self.detail_end_row,
            "trace_start_column": self.trace_start_column,
            "trace_headers": list(self.trace_headers),
        }


class CTableBaseService:
    """从 A 表完整复制并生成 C 表底稿。"""

    def build_base_table(self, a_file_path: Path, c_file_path: Path, a_schema: SheetFieldSchema) -> CTableBaseResult:
        """复制 A 表生成 C 表，并追加匹配追溯字段。"""

        # copy_workbook(...)：先完整复制 A 表，确保 C 表不是重新拼出来的空白表。
        excel_tool.copy_workbook(a_file_path, c_file_path)

        # load_workbook(...)：打开刚复制出的 C 表文件，后续只修改 C 表，不改 A 表。
        workbook = load_workbook(c_file_path)
        worksheet = excel_tool.get_sheet(workbook, a_schema.sheet_name or None)

        original_max_row = worksheet.max_row
        original_max_column = worksheet.max_column

        excel_tool.unmerge_cells_and_fill_values(worksheet)
        excel_tool.clear_sheet_fill(worksheet)
        excel_tool.format_date_column_as_month_day(
            worksheet,
            a_schema.date_field.column_index,
            a_schema.detail_start_row,
            a_schema.detail_end_row,
        )

        trace_start_column = original_max_column + 1
        self.write_trace_headers(worksheet, a_schema.header_row, trace_start_column)
        excel_tool.reset_view(worksheet, header_row=a_schema.header_row)
        excel_tool.auto_fit_columns(worksheet)

        # save(...)：把 C 表底稿写回任务目录。
        workbook.save(c_file_path)

        return CTableBaseResult(
            file_path=str(c_file_path),
            sheet_name=worksheet.title,
            original_max_row=original_max_row,
            original_max_column=original_max_column,
            header_row=a_schema.header_row,
            detail_start_row=a_schema.detail_start_row,
            detail_end_row=a_schema.detail_end_row,
            trace_start_column=trace_start_column,
            trace_headers=list(C_TRACE_HEADERS),
        )

    def write_trace_headers(self, worksheet: Worksheet, header_row: int, start_column: int) -> None:
        """在 C 表右侧写入固定追溯字段表头。"""

        for index, header in enumerate(C_TRACE_HEADERS):
            # cell(...)：按行列号定位单元格并写入表头文字。
            worksheet.cell(row=header_row, column=start_column + index, value=header)

        excel_tool.style_header(
            worksheet,
            row=header_row,
            start_col=start_column,
            end_col=start_column + len(C_TRACE_HEADERS) - 1,
        )
