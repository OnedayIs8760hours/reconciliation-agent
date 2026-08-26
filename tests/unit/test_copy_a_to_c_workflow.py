from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from agent.workflows.reconciliation.copy_a_to_c import copy_a_to_c_from_metadata


def test_copy_a_to_c_from_metadata_creates_c_draft_and_updates_json(tmp_path: Path) -> None:
    a_path = tmp_path / "A.xlsx"
    metadata_path = tmp_path / "metadata.json"

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "A表"
    worksheet["A1"] = "日期"
    worksheet["B1"] = "商品"
    worksheet["A2"] = date(2026, 7, 1)
    worksheet["B2"] = "商品A"
    worksheet["B2"].fill = PatternFill("solid", fgColor="FFFF00")
    worksheet["B3"] = "商品B"
    worksheet.merge_cells("A2:A3")
    workbook.save(a_path)

    metadata = {
        "task_id": "REC_TEST",
        "a_file_path": str(a_path),
        "b_file_path": str(tmp_path / "B.xlsx"),
        "reconcile_month": "2026-07",
        "a_preview": {
            "sheet_name": "A表",
            "rows": [
                {
                    "row_number": 1,
                    "cells": [
                        {"coordinate": "A1", "column": 1, "value": "日期"},
                        {"coordinate": "B1", "column": 2, "value": "商品"},
                    ],
                }
            ],
        },
        "llm_result": {"header_row_guess": 1},
        "product_mapping": {"structure": {"a_sheet": {"header_row_guess": 1}}},
    }
    metadata_path.write_text(json.dumps(metadata, ensure_ascii=False), encoding="utf-8")

    c_path = copy_a_to_c_from_metadata(metadata_path)

    c_workbook = load_workbook(c_path)
    c_sheet = c_workbook["A表"]
    assert c_path == tmp_path / "C.xlsx"
    assert not list(c_sheet.merged_cells.ranges)
    assert c_sheet["A1"].value == "七月"
    assert c_sheet["A2"].value == "7月1日"
    assert c_sheet["A3"].value == "7月1日"
    assert c_sheet["B2"].fill.fill_type is None

    updated_metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    assert updated_metadata["c_file_path"] == str(c_path)
    assert updated_metadata["workflow_steps"][0]["step"] == "01_copy_a_to_c"
    assert updated_metadata["workflow_steps"][0]["date_column"] == "A"
