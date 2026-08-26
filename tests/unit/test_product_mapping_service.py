from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook

from agent.llm_providers.base import LLMResponse
from backend.services.product_mapping_service import ProductMappingService


class FakeProductMappingAgent:
    """测试用假 LLM，只返回固定 JSON，避免单元测试调用真实模型。"""

    def analyze_sheet_structure(self, a_preview: object, b_preview: object) -> LLMResponse:
        """模拟 LLM 根据表结构识别商品字段。"""

        text = """
        {
          "a_sheet": {
            "header_row_guess": 2,
            "field_name": "产品名称",
            "confidence": 0.95,
            "reason": "A 表样例中产品名称最像商品字段"
          },
          "b_sheet": {
            "header_row_guess": 1,
            "field_name": "规格",
            "confidence": 0.96,
            "reason": "B 表样例中规格最像商品字段"
          }
        }
        """
        return LLMResponse(text=text, model="fake-model", provider="deepseek")

    def analyze_product_mapping(self, a_products: list[str], b_products: list[str]) -> LLMResponse:
        """模拟 LLM 生成商品映射结果。"""

        text = """
        {
          "normalization_rules": {
            "b_prefix_to_ignore": [],
            "ignorable_separators": []
          },
          "mappings": [
            {
              "standard": "商品A",
              "a_value": "商品A",
              "b_value": "商品A",
              "confidence": 1.0
            }
          ],
          "unmatched_a": ["商品B"],
          "unmatched_b": ["商品C"],
          "need_review": []
        }
        """
        return LLMResponse(text=text, model="fake-model", provider="deepseek")


def create_a_file(file_path: Path) -> None:
    """创建测试 A 表，表头故意放在第 2 行，验证代码不会写死第 1 行。"""

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "A表"
    worksheet.merge_cells("A1:C1")
    worksheet["A1"] = "A 表标题"
    worksheet["A2"] = "日期"
    worksheet["B2"] = "产品名称"
    worksheet["C2"] = "数量"
    worksheet["A3"] = "2026-07-01"
    worksheet["B3"] = " 商品A "
    worksheet["C3"] = 1
    worksheet["A4"] = "2026-07-02"
    worksheet["B4"] = "商品B"
    worksheet["C4"] = 2
    # save(...)：保存为真实 xlsx 文件，测试读取 Excel 的完整流程。
    workbook.save(file_path)


def create_b_file(file_path: Path) -> None:
    """创建测试 B 表，字段名与 A 表不同，验证字段由 LLM 返回。"""

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "B表"
    worksheet["A1"] = "规格"
    worksheet["B1"] = "数量"
    worksheet["A2"] = "商品A"
    worksheet["B2"] = 1
    worksheet["A3"] = "商品C"
    worksheet["B3"] = 3
    workbook.save(file_path)


def test_product_mapping_service_uses_llm_guessed_fields(tmp_path: Path) -> None:
    """商品映射服务应该使用 LLM 识别出的字段，不应该写死 型号/规格 或固定表头行。"""

    a_file_path = tmp_path / "a.xlsx"
    b_file_path = tmp_path / "b.xlsx"
    create_a_file(a_file_path)
    create_b_file(b_file_path)

    service = ProductMappingService(llm_agent=FakeProductMappingAgent())  # type: ignore[arg-type]

    result = service.build_product_mapping(a_file_path, b_file_path)

    assert result["a_column_name"] == "产品名称"
    assert result["b_column_name"] == "规格"
    assert result["a_unique"] == ["商品A", "商品B"]
    assert result["b_unique"] == ["商品A", "商品C"]
    assert result["a_b_intersection"] == {
        "matched": [{"a": "商品A", "b": "商品A"}],
        "a_unmatched": ["商品B"],
        "b_unused": ["商品C"],
    }
    mapping_item = result["result"]["mappings"][0]  # type: ignore[index]
    assert "reason" not in mapping_item
    assert result["summary"]["mapping_count"] == 1  # type: ignore[index]


def test_product_mapping_service_reports_empty_llm_field(tmp_path: Path) -> None:
    """如果 LLM 没有识别出字段，服务应该给出清晰错误，而不是继续用空字段读取。"""

    class EmptyFieldAgent(FakeProductMappingAgent):
        def analyze_sheet_structure(self, a_preview: object, b_preview: object) -> LLMResponse:
            """模拟 LLM 未识别出 A 表字段。"""

            text = """
            {
              "a_sheet": {"header_row_guess": 1, "field_name": "", "confidence": 0.1, "reason": "看不出"},
              "b_sheet": {"header_row_guess": 1, "field_name": "规格", "confidence": 0.9, "reason": "可识别"}
            }
            """
            return LLMResponse(text=text, model="fake-model", provider="deepseek")

    a_file_path = tmp_path / "a.xlsx"
    b_file_path = tmp_path / "b.xlsx"
    create_a_file(a_file_path)
    create_b_file(b_file_path)

    service = ProductMappingService(llm_agent=EmptyFieldAgent())  # type: ignore[arg-type]

    try:
        service.build_product_mapping(a_file_path, b_file_path)
    except ValueError as exc:
        assert "LLM 没有识别出 A 表商品字段" in str(exc)
    else:
        raise AssertionError("LLM 字段为空时应该抛出 ValueError")


def test_product_mapping_service_does_not_match_identical_values_without_llm_mapping(
    tmp_path: Path,
) -> None:
    """a_b_intersection 应只来自 LLM 返回的 mappings，不应自行补完全匹配。"""

    class EmptyMappingAgent(FakeProductMappingAgent):
        def analyze_product_mapping(self, a_products: list[str], b_products: list[str]) -> LLMResponse:
            text = """
            {
              "normalization_rules": {},
              "mappings": [],
              "unmatched_a": [],
              "unmatched_b": [],
              "need_review": []
            }
            """
            return LLMResponse(text=text, model="fake-model", provider="deepseek")

    a_file_path = tmp_path / "a.xlsx"
    b_file_path = tmp_path / "b.xlsx"
    create_a_file(a_file_path)
    create_b_file(b_file_path)

    service = ProductMappingService(llm_agent=EmptyMappingAgent())  # type: ignore[arg-type]

    result = service.build_product_mapping(a_file_path, b_file_path)

    assert result["a_b_intersection"] == {
        "matched": [],
        "a_unmatched": ["商品A", "商品B"],
        "b_unused": ["商品A", "商品C"],
    }


def test_product_mapping_service_does_not_fallback_when_llm_returns_empty_text(
    tmp_path: Path,
) -> None:
    """商品映射 LLM 空返回时，不应写死规则或本地生成 mappings。"""

    class EmptyTextMappingAgent(FakeProductMappingAgent):
        def analyze_sheet_structure(self, a_preview: object, b_preview: object) -> LLMResponse:
            text = """
            {
              "a_sheet": {"header_row_guess": 1, "field_name": "产品名称", "confidence": 0.95, "reason": "可识别"},
              "b_sheet": {"header_row_guess": 1, "field_name": "规格", "confidence": 0.95, "reason": "可识别"}
            }
            """
            return LLMResponse(text=text, model="fake-model", provider="deepseek")

        def analyze_product_mapping(self, a_products: list[str], b_products: list[str]) -> LLMResponse:
            return LLMResponse(text="", model="fake-model", provider="deepseek")

    a_file_path = tmp_path / "a.xlsx"
    b_file_path = tmp_path / "b.xlsx"

    a_workbook = Workbook()
    a_sheet = a_workbook.active
    a_sheet.title = "A表"
    a_sheet["A1"] = "产品名称"
    a_sheet["A2"] = "XN6012-奶油色-带字款"
    a_sheet["A3"] = "WK9648-红箱-毛衣狗-M-鹿角"
    a_sheet["A4"] = "WK9648-收纳箱-红箱-小熊-M圆角"
    a_workbook.save(a_file_path)

    b_workbook = Workbook()
    b_sheet = b_workbook.active
    b_sheet.title = "B表"
    b_sheet["A1"] = "规格"
    b_sheet["A2"] = "XN6012-奶油色-带字款"
    b_sheet["A3"] = "WK9648-收纳箱-红箱-毛衣狗-M-鹿角"
    b_sheet["A4"] = "WK9648-收纳箱-红箱-毛衣狗-M-圆角"
    b_workbook.save(b_file_path)

    service = ProductMappingService(llm_agent=EmptyTextMappingAgent())  # type: ignore[arg-type]

    result = service.build_product_mapping(a_file_path, b_file_path)

    assert result["a_b_intersection"] == {
        "matched": [],
        "a_unmatched": [
            "XN6012-奶油色-带字款",
            "WK9648-红箱-毛衣狗-M-鹿角",
            "WK9648-收纳箱-红箱-小熊-M圆角",
        ],
        "b_unused": [
            "XN6012-奶油色-带字款",
            "WK9648-收纳箱-红箱-毛衣狗-M-鹿角",
            "WK9648-收纳箱-红箱-毛衣狗-M-圆角",
        ],
    }
    assert result["result"]["normalization_rules"] == {}  # type: ignore[index]
    assert result["result"]["mappings"] == []  # type: ignore[index]
    assert result["summary"]["mapping_count"] == 0  # type: ignore[index]
    assert "第 1 批" in result["result"]["parse_error"]  # type: ignore[index]
    assert "LLM 返回内容不是合法 JSON" in result["result"]["parse_error"]  # type: ignore[index]


def test_product_mapping_service_batches_a_products_by_twenty() -> None:
    """商品映射应按 A 表每 20 个一批调用 LLM，并合并每批 mappings。"""

    class BatchAgent:
        provider = "deepseek"
        model = "fake-model"

        def __init__(self) -> None:
            self.a_batches: list[list[str]] = []
            self.b_batches: list[list[str]] = []

        def analyze_product_mapping(self, a_products: list[str], b_products: list[str]) -> LLMResponse:
            self.a_batches.append(list(a_products))
            self.b_batches.append(list(b_products))
            mappings = [
                {
                    "a_value": a_value,
                    "b_value": a_value,
                    "confidence": 1.0,
                }
                for a_value in a_products
                if a_value in b_products
            ]
            text = (
                "{"
                "\"mappings\": "
                f"{json.dumps(mappings, ensure_ascii=False)}, "
                "\"need_review\": []"
                "}"
            )
            return LLMResponse(text=text, model=self.model, provider=self.provider)

    a_unique = [f"商品{i:02d}" for i in range(45)]
    b_unique = list(a_unique)
    agent = BatchAgent()
    service = ProductMappingService(llm_agent=agent)  # type: ignore[arg-type]

    mapping_result, _, _ = service.analyze_product_mapping_in_batches(a_unique, b_unique)

    assert [len(batch) for batch in agent.a_batches] == [20, 20, 5]
    assert [len(batch) for batch in agent.b_batches] == [45, 25, 5]
    assert mapping_result.mapping_count == 45
    assert mapping_result.unmatched_a == []
    assert mapping_result.unmatched_b == []
    assert "reason" not in mapping_result.mappings[0].to_dict()
