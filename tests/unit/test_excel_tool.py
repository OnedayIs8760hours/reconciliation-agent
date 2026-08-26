from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from tools.excel_tool import excel_tool


def test_read_sheet_preview_handles_merged_cells(tmp_path: Path) -> None:
    """读取含合并单元格的表格时，不应该因为 MergedCell 没有 is_date 而报错。"""

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "测试表"

    # merge_cells(...)：模拟用户真实表格里的合并标题单元格。
    worksheet.merge_cells("A1:C1")
    worksheet["A1"] = "好一家账单"
    worksheet["A2"] = "日期"
    worksheet["B2"] = "产品名称"
    worksheet["C2"] = "数量"
    worksheet["A3"] = "2026-07-01"
    worksheet["B3"] = "商品A"
    worksheet["C3"] = 10

    file_path = tmp_path / "merged.xlsx"
    # save(...)：把测试工作簿保存成真实 xlsx，确保读取逻辑走 openpyxl 文件解析。
    workbook.save(file_path)

    preview = excel_tool.read_sheet_preview(file_path, rows=3, max_columns=3)

    assert preview.sheet_name == "测试表"
    assert preview.rows[0].cells[0].display_text == "好一家账单"
    assert preview.rows[1].cells[1].display_text == "产品名称"


def test_clean_product_text_preserves_business_content() -> None:
    """商品清洗只处理格式，不删除型号、颜色、款、新、星号等业务内容。"""

    value = "　新款 １２３*红色\n"

    cleaned_value = excel_tool.clean_product_text(value)

    assert cleaned_value == "新款 123*红色"


def test_unique_product_values_keeps_first_order() -> None:
    """商品去重应该去掉空值和重复值，并保留第一次出现的顺序。"""

    values = [" 商品A ", "", None, "商品B", "商品A", "商品Ｃ", "商品C"]

    result = excel_tool.unique_product_values(values)

    assert result == ["商品A", "商品B", "商品C"]


def test_get_unique_column_values_only_deduplicates(tmp_path: Path) -> None:
    """读取唯一列值时只去重，不做清洗、去空或全角转换。"""

    workbook = Workbook()
    worksheet = workbook.active
    worksheet["A1"] = "产品名称"
    worksheet["A2"] = " 商品A "
    worksheet["A3"] = ""
    worksheet["A4"] = None
    worksheet["A5"] = "商品A"
    worksheet["A6"] = " 商品A "
    worksheet["A7"] = "商品Ｃ"
    worksheet["A8"] = "商品C"

    file_path = tmp_path / "unique.xlsx"
    workbook.save(file_path)

    result = excel_tool.get_unique_column_values(file_path, "产品名称")

    assert result == [" 商品A ", None, "商品A", "商品Ｃ", "商品C"]
