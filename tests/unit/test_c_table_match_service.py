from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook

from backend.domain.match_result import MatchSummary
from backend.domain.reconciliation_record import ATableRecord, BTableRecord
from backend.domain.reconciliation_schema import FieldLocation, SheetFieldSchema
from backend.services.b_missing_marker_service import BMissingMarkerResult
from backend.services.c_table_base_service import CTableBaseResult
from backend.services.c_table_match_service import CTableMatchService
from backend.services.product_mapping_service import ProductMappingLookup
from backend.services.reconciliation_verify_service import ReconciliationVerifyService
from backend.services.reverse_verify_service import ReverseVerifySummary
from backend.domain.product_mapping import ProductMappingItem, ProductMappingResult


def build_b_record(source_row: int, quantity: float, amount: float, spec: str = "B表规格A") -> BTableRecord:
    """构造测试用 B 表记录，避免每个测试重复填写字段。"""

    return BTableRecord(
        source_row=source_row,
        direction="入库",
        warehouse="测试仓",
        system_time=date(2026, 7, source_row),
        document_no=f"DOC-{source_row}",
        sku=f"SKU-{source_row}",
        product_name="测试商品",
        spec=spec,
        quantity=quantity,
        unit_price=10,
        amount=amount,
        is_current_month=True,
    )


def test_find_total_match_supports_non_consecutive_b_records() -> None:
    """多行匹配应该能找到非连续 B 表记录组合。"""

    service = CTableMatchService()
    a_record = ATableRecord(
        source_row=2,
        date_value=date(2026, 7, 1),
        product_value="A表商品A",
        quantity=8,
        unit_price=10,
        amount=80,
    )
    candidates = [
        build_b_record(2, 3, 30),
        build_b_record(3, 4, 40),
        build_b_record(4, 5, 50),
    ]

    matched_records = service.find_total_match(a_record, candidates)

    assert [record.source_row for record in matched_records] == [2, 4]


def test_find_multi_match_uses_product_mapping_before_combining() -> None:
    """多行匹配组合前应该先按商品映射筛选候选 B 记录。"""

    service = CTableMatchService()
    a_record = ATableRecord(
        source_row=2,
        date_value=date(2026, 7, 1),
        product_value="A表商品A",
        quantity=8,
        unit_price=10,
        amount=80,
    )
    mapping_result = ProductMappingResult(
        mappings=[
            ProductMappingItem(
                standard="标准商品A",
                a_value="A表商品A",
                b_value="B表规格A",
            )
        ]
    )
    product_lookup = ProductMappingLookup(mapping_result)
    candidates = [
        build_b_record(2, 3, 30, spec="B表规格A"),
        build_b_record(3, 99, 990, spec="不相关规格"),
        build_b_record(4, 5, 50, spec="B表规格A"),
    ]

    matched_records = service.find_multi_match(a_record, candidates, product_lookup)

    assert [record.source_row for record in matched_records] == [2, 4]


def test_scan_formula_errors_uses_named_b_marked_sheet(tmp_path: Path) -> None:
    """公式错误扫描应该使用 LLM 识别出的 B 表工作表，不能默认 active。"""

    c_file_path = tmp_path / "C.xlsx"
    b_marked_file_path = tmp_path / "B_marked.xlsx"

    c_workbook = Workbook()
    c_worksheet = c_workbook.active
    c_worksheet.title = "C表"
    c_worksheet["A1"] = "正常"
    c_workbook.save(c_file_path)

    b_workbook = Workbook()
    active_sheet = b_workbook.active
    active_sheet.title = "说明页"
    active_sheet["A1"] = "#REF!"
    real_sheet = b_workbook.create_sheet("B明细")
    real_sheet["A1"] = "正常"
    b_workbook.save(b_marked_file_path)

    service = ReconciliationVerifyService()
    errors = service.scan_formula_errors(c_file_path, b_marked_file_path, "C表", "B明细")

    assert errors == []


def test_verify_report_keeps_b_sheet_name_in_marker_result(tmp_path: Path) -> None:
    """验收报告应该能通过标注结果中的工作表名扫描 B 标注表。"""

    a_file_path = tmp_path / "A.xlsx"
    c_file_path = tmp_path / "C.xlsx"
    b_marked_file_path = tmp_path / "B_marked.xlsx"
    report_file_path = tmp_path / "verify_report.json"

    for file_path in [a_file_path, c_file_path]:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "A明细"
        worksheet.cell(row=1, column=1, value="数量")
        worksheet.cell(row=1, column=2, value="金额")
        worksheet.cell(row=2, column=1, value=5)
        worksheet.cell(row=2, column=2, value=50)
        workbook.save(file_path)

    b_workbook = Workbook()
    b_workbook.active.title = "说明页"
    b_sheet = b_workbook.create_sheet("B明细")
    b_sheet.cell(row=1, column=1, value="正常")
    b_workbook.save(b_marked_file_path)

    c_table_result = CTableBaseResult(
        file_path=str(c_file_path),
        sheet_name="A明细",
        original_max_row=2,
        original_max_column=2,
        header_row=1,
        detail_start_row=2,
        detail_end_row=2,
        trace_start_column=3,
        trace_headers=[],
    )
    marker_result = BMissingMarkerResult(
        file_path=str(b_marked_file_path),
        marked_count=0,
        marker_start_column=2,
        marker_headers=[],
        sheet_name="B明细",
    )

    service = ReconciliationVerifyService()
    report = service.verify(
        a_file_path,
        c_file_path,
        b_marked_file_path,
        c_table_result,
        MatchSummary(),
        ReverseVerifySummary(),
        marker_result,
        quantity_column=1,
        amount_column=2,
        report_file_path=report_file_path,
    )

    assert report.passed is True
    assert report.formula_errors == []
    assert report_file_path.exists()
