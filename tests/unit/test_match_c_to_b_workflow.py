from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

from agent.workflows.reconciliation.match_c_to_b import (
    infer_c_columns,
    match_c_to_b_from_metadata,
    parse_month_day,
)


def test_match_c_to_b_from_metadata_expands_multiple_b_rows(tmp_path: Path) -> None:
    c_path = tmp_path / "C.xlsx"
    b_path = tmp_path / "B.xlsx"
    metadata_path = tmp_path / "metadata.json"

    c_workbook = Workbook()
    c_sheet = c_workbook.active
    c_sheet["A1"] = "7月份"
    c_sheet["A2"] = "7月1日"
    c_sheet["B2"] = "AY921-猫咪垃圾桶-红黄奶*"
    c_sheet["C2"] = 462
    c_sheet["D2"] = 20.86
    c_sheet["E2"] = 9009
    c_sheet["A3"] = "合计"
    c_sheet["E3"] = "=SUM(E2:E2)"
    c_workbook.save(c_path)

    b_workbook = Workbook()
    b_sheet = b_workbook.active
    headers = [
        "仓库",
        "系统出入库时间",
        "单据编号",
        "货品编号",
        "货品名称",
        "规格",
        "入库数量",
        "入库成本单价",
        "入库成本金额",
        "出库数量",
        "出库成本单价",
        "出库成本金额",
    ]
    b_sheet.append(headers)
    b_sheet.append(
        [
            "仓库A",
            "2026-07-01 11:19:04",
            "CRK1",
            "M-AY-92",
            "猫咪垃圾桶",
            "AY921-猫咪垃圾桶-红黄奶",
            449,
            19.5,
            8755.5,
            None,
            None,
            None,
        ]
    )
    b_sheet.append(
        [
            "仓库A",
            "2026-07-01 11:19:11",
            "CRK2",
            "M-AY-92",
            "猫咪垃圾桶",
            "AY921-猫咪垃圾桶-红黄奶",
            13,
            19.5,
            253.5,
            None,
            None,
            None,
        ]
    )
    b_workbook.save(b_path)

    metadata = {
        "task_id": "REC_TEST",
        "c_file_path": str(c_path),
        "b_file_path": str(b_path),
        "product_mapping": {
            "result": {
                "mappings": [
                    {
                        "a_value": "AY921-猫咪垃圾桶-红黄奶*",
                        "b_value": "AY921-猫咪垃圾桶-红黄奶",
                    }
                ]
            }
        },
    }
    metadata_path.write_text(json.dumps(metadata, ensure_ascii=False), encoding="utf-8")

    match_c_to_b_from_metadata(metadata_path)

    workbook = load_workbook(c_path, data_only=False)
    sheet = workbook.active

    assert sheet.max_row == 4
    assert sheet["G1"].value == "匹配状态"
    assert sheet["G2"].value == "已匹配/多行展开"
    assert sheet["I2"].value == 2
    assert sheet["J2"].value == "入库"
    assert sheet["K2"].value == 0
    assert sheet["L2"].value == "仓库A"
    assert sheet["R2"].value == 449
    assert sheet["G3"].value == "已匹配/多行展开"
    assert sheet["A3"].value is None
    assert sheet["B3"].value is None
    assert sheet["I3"].value == 3
    assert sheet["R3"].value == 13
    assert sheet["A4"].value == "合计"
    assert sheet["E4"].value == "=SUM(E2:E3)"

    updated_metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    assert updated_metadata["workflow_steps"][0]["step"] == "02_match_c_to_b"
    assert updated_metadata["workflow_steps"][0]["summary"] == {
        "source_rows": 1,
        "matched_source_rows": 1,
        "expanded_b_rows": 2,
        "unmatched_source_rows": 0,
        "review_source_rows": 0,
        "appended_b_rows": 0,
    }

    match_c_to_b_from_metadata(metadata_path)

    workbook_after_second_run = load_workbook(c_path, data_only=False)
    assert workbook_after_second_run.active.max_row == 4


def test_match_c_to_b_appends_headers_to_detected_header_row(tmp_path: Path) -> None:
    c_path = tmp_path / "C.xlsx"
    b_path = tmp_path / "B.xlsx"
    metadata_path = tmp_path / "metadata.json"

    c_workbook = Workbook()
    c_sheet = c_workbook.active
    c_sheet["A1"] = "2026年7月对账单"
    c_sheet["A2"] = "客户"
    c_sheet["B2"] = "日期"
    c_sheet["C2"] = "型号"
    c_sheet["D2"] = "数量"
    c_sheet["E2"] = "单价"
    c_sheet["F2"] = "金额"
    c_sheet["A3"] = "玖鸣"
    c_sheet["B3"] = "7月1日"
    c_sheet["C3"] = "规格A"
    c_sheet["D3"] = 10
    c_sheet["E3"] = 2
    c_sheet["F3"] = "=D3*E3"
    c_workbook.save(c_path)

    b_workbook = Workbook()
    b_sheet = b_workbook.active
    b_sheet.append(
        [
            "仓库",
            "系统出入库时间",
            "单据编号",
            "货品编号",
            "货品名称",
            "规格",
            "入库数量",
            "入库成本单价",
            "入库成本金额",
            "出库数量",
            "出库成本单价",
            "出库成本金额",
        ]
    )
    b_sheet.append(["仓库A", "2026-07-01", "CRK1", "M1", "商品A", "规格A", 10, 2, 20, None, None, None])
    b_workbook.save(b_path)

    metadata = {
        "task_id": "REC_TEST",
        "c_file_path": str(c_path),
        "b_file_path": str(b_path),
        "product_mapping": {"result": {"mappings": [{"a_value": "规格A", "b_value": "规格A"}]}},
    }
    metadata_path.write_text(json.dumps(metadata, ensure_ascii=False), encoding="utf-8")

    match_c_to_b_from_metadata(metadata_path)

    workbook = load_workbook(c_path, data_only=False)
    sheet = workbook.active
    assert sheet["G1"].value is None
    assert sheet["G2"].value == "匹配状态"
    assert sheet["G3"].value == "已匹配"


def test_match_c_to_b_allows_one_a_value_to_match_multiple_b_specs(tmp_path: Path) -> None:
    c_path = tmp_path / "C.xlsx"
    b_path = tmp_path / "B.xlsx"
    metadata_path = tmp_path / "metadata.json"

    c_workbook = Workbook()
    c_sheet = c_workbook.active
    c_sheet.append(["日期", "规格", "数量", "单价", "金额"])
    c_sheet.append(["7月1日", "规格A", 10, 2, 20])
    c_sheet.append(["7月1日", "规格A", 12, 2, 24])
    c_workbook.save(c_path)

    b_workbook = Workbook()
    b_sheet = b_workbook.active
    b_sheet.append(
        [
            "仓库",
            "系统出入库时间",
            "单据编号",
            "货品编号",
            "货品名称",
            "规格",
            "入库数量",
            "入库成本单价",
            "入库成本金额",
            "出库数量",
            "出库成本单价",
            "出库成本金额",
        ]
    )
    b_sheet.append(["仓库A", "2026-07-01", "CRK1", "M1", "商品A", "B规格1", 10, 2, 20, None, None, None])
    b_sheet.append(["仓库A", "2026-07-01", "CRK2", "M2", "商品B", "B规格2", 12, 2, 24, None, None, None])
    b_workbook.save(b_path)

    metadata = {
        "c_file_path": str(c_path),
        "b_file_path": str(b_path),
        "product_mapping": {
            "result": {
                "mappings": [
                    {"a_value": "规格A", "b_value": "B规格1"},
                    {"a_value": "规格A", "b_value": "B规格2"},
                ]
            }
        },
    }
    metadata_path.write_text(json.dumps(metadata, ensure_ascii=False), encoding="utf-8")

    match_c_to_b_from_metadata(metadata_path)

    workbook = load_workbook(c_path, data_only=False)
    sheet = workbook.active
    assert sheet["G2"].value == "已匹配"
    assert sheet["I2"].value == 2
    assert sheet["G3"].value == "已匹配"
    assert sheet["I3"].value == 3


def test_match_c_to_b_marks_date_review_when_only_date_differs(tmp_path: Path) -> None:
    c_path = tmp_path / "C.xlsx"
    b_path = tmp_path / "B.xlsx"
    metadata_path = tmp_path / "metadata.json"

    c_workbook = Workbook()
    c_sheet = c_workbook.active
    c_sheet.append(["日期", "规格", "数量", "单价", "金额"])
    c_sheet.append(["7月1日", "规格A", 10, 2, 20])
    c_workbook.save(c_path)

    b_workbook = Workbook()
    b_sheet = b_workbook.active
    b_sheet.append(
        [
            "仓库",
            "系统出入库时间",
            "单据编号",
            "货品编号",
            "货品名称",
            "规格",
            "入库数量",
            "入库成本单价",
            "入库成本金额",
            "出库数量",
            "出库成本单价",
            "出库成本金额",
        ]
    )
    b_sheet.append(["仓库A", "2026-07-02", "CRK1", "M1", "商品A", "B规格", 10, 2, 20, None, None, None])
    b_workbook.save(b_path)

    metadata = {
        "c_file_path": str(c_path),
        "b_file_path": str(b_path),
        "product_mapping": {"result": {"mappings": [{"a_value": "规格A", "b_value": "B规格"}]}},
    }
    metadata_path.write_text(json.dumps(metadata, ensure_ascii=False), encoding="utf-8")

    match_c_to_b_from_metadata(metadata_path)

    workbook = load_workbook(c_path, data_only=False)
    sheet = workbook.active
    assert sheet["G2"].value == "已匹配/日期需复核"
    assert sheet["I2"].value == 2


def test_match_c_to_b_appends_unmatched_b_records_before_total(tmp_path: Path) -> None:
    c_path = tmp_path / "C.xlsx"
    b_path = tmp_path / "B.xlsx"
    metadata_path = tmp_path / "metadata.json"

    c_workbook = Workbook()
    c_sheet = c_workbook.active
    c_sheet.append(["日期", "规格", "数量", "单价", "金额"])
    c_sheet.append(["7月1日", "规格A", 10, 2, 20])
    c_sheet.append(["合计", None, None, None, "=SUM(E2:E2)"])
    c_workbook.save(c_path)

    b_workbook = Workbook()
    b_sheet = b_workbook.active
    b_sheet.append(
        [
            "仓库",
            "系统出入库时间",
            "单据编号",
            "货品编号",
            "货品名称",
            "规格",
            "入库数量",
            "入库成本单价",
            "入库成本金额",
            "出库数量",
            "出库成本单价",
            "出库成本金额",
        ]
    )
    b_sheet.append(["仓库A", "2026-07-01", "CRK1", "M1", "商品A", "B规格1", 10, 2, 20, None, None, None])
    b_sheet.append(["仓库A", "2026-07-03", "CRK2", "M2", "商品B", "B规格2", 5, 3, 15, None, None, None])
    b_workbook.save(b_path)

    metadata = {
        "c_file_path": str(c_path),
        "b_file_path": str(b_path),
        "product_mapping": {"result": {"mappings": [{"a_value": "规格A", "b_value": "B规格1"}]}},
    }
    metadata_path.write_text(json.dumps(metadata, ensure_ascii=False), encoding="utf-8")

    match_c_to_b_from_metadata(metadata_path)

    workbook = load_workbook(c_path, data_only=False)
    sheet = workbook.active
    assert sheet["G2"].value == "已匹配"
    assert sheet["G3"].value == "B表未匹配追加"
    assert sheet["B3"].value == "B规格2"
    assert sheet["C3"].value == 5
    assert sheet["I3"].value == 3
    assert sheet["A4"].value == "合计"
    assert sheet["E4"].value == "=SUM(E2:E3)"

    updated_metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    assert updated_metadata["workflow_steps"][0]["summary"]["appended_b_rows"] == 1


def test_parse_month_day_handles_numeric_month_day() -> None:
    assert parse_month_day(5.5) == (5, 5)
    assert parse_month_day(6.09) == (6, 9)


def test_infer_c_columns_uses_date_values_when_header_is_month_title() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["五月", "单据类型", "单据编号", "货品编号", "货品名称", "规格", "入库数量", "单价", "金额"])
    sheet.append(["2026-05-09 09:56:19", "采购入库", "CRK1", "M1", "商品A", "规格A", 10, 2, 20])

    columns = infer_c_columns(sheet)

    assert columns.date_col == 1
    assert columns.spec_col == 6
    assert columns.qty_col == 7
    assert columns.price_col == 8
    assert columns.amount_col == 9
    assert columns.first_detail_row == 2
