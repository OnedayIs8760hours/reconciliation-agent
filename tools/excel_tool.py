"""Excel 读写与表格数据处理工具。

这个模块统一封装 openpyxl 的常用操作，业务层只需要关注“工作簿、工作表、行数据、表头”这些概念，
避免在对账流程里反复编写 Excel 文件读写细节。
"""

from __future__ import annotations

from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any
import unicodedata

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

ExcelValue = str | int | float | bool | date | datetime | time | None
RowValues = Sequence[ExcelValue]
Record = Mapping[str, ExcelValue]


@dataclass(frozen=True)
class ExcelSheetData:
    """从 Excel 工作表读取出的结构化数据。"""

    sheet_name: str
    headers: list[str]
    rows: list[dict[str, ExcelValue]]


@dataclass(frozen=True)
class ExcelPreviewCell:
    """提供给 LLM 的单元格预览数据，保留值、类型和 Excel 显示格式。"""

    coordinate: str
    column: int
    value: ExcelValue
    python_type: str
    excel_data_type: str
    number_format: str
    is_date: bool
    display_text: str


@dataclass(frozen=True)
class ExcelPreviewRow:
    """提供给 LLM 观察 Excel 表格结构的一行预览数据。"""

    row_number: int
    cells: list[ExcelPreviewCell]


@dataclass(frozen=True)
class ExcelSheetPreview:
    """提供给 LLM 的 Excel 工作表预览信息。"""

    file_path: str
    sheet_name: str
    max_row: int
    max_column: int
    rows: list[ExcelPreviewRow]


class ExcelTool:
    """Excel 文件操作工具类。

    主要能力：
    - 打开/保存工作簿
    - 按表头读取工作表数据
    - 创建工作表并写入二维数据或字典数据
    - 设置冻结窗格、筛选、列宽和表头样式
    """

    DEFAULT_HEADER_FILL = "D9EAF7"
    DEFAULT_HEADER_FONT_COLOR = "000000"

    def load_workbook(self, file_path: str | Path, *, data_only: bool = True) -> Workbook:
        """读取 Excel 工作簿。

        Args:
            file_path: .xlsx 文件路径。
            data_only: 是否读取公式缓存值。True 更适合对账场景，False 可读取公式本身。
        """

        path = self._ensure_xlsx_path(file_path, must_exist=True)
        return load_workbook(path, data_only=data_only)

    def create_workbook(self, *, default_sheet_name: str | None = None) -> Workbook:
        """创建一个新的 Excel 工作簿。"""

        workbook = Workbook()
        if default_sheet_name:
            workbook.active.title = default_sheet_name
        return workbook

    def save_workbook(self, workbook: Workbook, file_path: str | Path) -> Path:
        """保存工作簿，如果父目录不存在会自动创建。"""

        path = self._ensure_xlsx_path(file_path, must_exist=False)
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
        return path

    def get_sheet(self, workbook: Workbook, sheet_name: str | None = None) -> Worksheet:
        """获取指定工作表；不传 sheet_name 时返回当前活动工作表。"""

        if sheet_name is None:
            return workbook.active

        if sheet_name not in workbook.sheetnames:
            available = ", ".join(workbook.sheetnames)
            raise ValueError(f"工作表不存在：{sheet_name}；可用工作表：{available}")

        return workbook[sheet_name]

    def read_sheet_preview(
        self,
        file_path: str | Path,
        *,
        sheet_name: str | None = None,
        rows: int = 8,
        start_row: int = 1,
        max_columns: int | None = None,
        data_only: bool = True,
        skip_empty_tail_columns: bool = True,
    ) -> ExcelSheetPreview:
        """读取工作表前几行，方便 LLM 判断 Excel 数据长什么样。

        这个方法不假设哪一行是表头，只按原始行号返回样例数据，并保留单元格值、
        Python 类型、Excel 数据类型和数字格式，适合传给 LLM 做：
        - 判断表头在哪一行
        - 判断字段含义
        - 判断 A/B 表结构差异
        - 生成后续读取参数
        """

        if rows < 1:
            raise ValueError("rows 必须大于等于 1")
        if start_row < 1:
            raise ValueError("start_row 必须大于等于 1")
        if max_columns is not None and max_columns < 1:
            raise ValueError("max_columns 必须大于等于 1")

        workbook = self.load_workbook(file_path, data_only=data_only)
        worksheet = self.get_sheet(workbook, sheet_name)
        end_row = min(worksheet.max_row, start_row + rows - 1)
        end_col = min(worksheet.max_column, max_columns) if max_columns else worksheet.max_column

        preview_rows: list[ExcelPreviewRow] = []
        for row_number, row in enumerate(
            worksheet.iter_rows(
                min_row=start_row,
                max_row=end_row,
                min_col=1,
                max_col=end_col,
                values_only=False,
            ),
            start=start_row,
        ):
            cells = [self.build_preview_cell(cell) for cell in row]
            if skip_empty_tail_columns:
                cells = self.trim_trailing_empty_cells(cells)
            preview_rows.append(ExcelPreviewRow(row_number=row_number, cells=cells))

        return ExcelSheetPreview(
            file_path=str(file_path),
            sheet_name=worksheet.title,
            max_row=worksheet.max_row,
            max_column=worksheet.max_column,
            rows=preview_rows,
        )

    def read_sheet_as_records(
        self,
        file_path: str | Path,
        *,
        sheet_name: str | None = None,
        header_row: int = 1,
        data_only: bool = True,
        skip_empty_rows: bool = True,
    ) -> ExcelSheetData:
        """按表头把工作表读取为字典列表。

        返回的每一行形如：{"供应商": "A公司", "金额": 100.0}。
        空表头列会被忽略；重复表头会自动追加序号，避免字段覆盖。
        """

        workbook = self.load_workbook(file_path, data_only=data_only)
        worksheet = self.get_sheet(workbook, sheet_name)
        headers = self.read_headers(worksheet, header_row=header_row)

        rows: list[dict[str, ExcelValue]] = []
        for row in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
            values = list(row[: len(headers)])
            if skip_empty_rows and self.is_empty_row(values):
                continue

            record = {
                header: self.normalize_cell_value(values[index] if index < len(values) else None)
                for index, header in enumerate(headers)
                if header
            }
            rows.append(record)

        return ExcelSheetData(sheet_name=worksheet.title, headers=headers, rows=rows)

    def get_column_values(
        self,
        file_path: str | Path,
        column_name: str,
        *,
        sheet_name: str | None = None,
        header_row: int = 1,
        data_only: bool = True,
    ) -> list[ExcelValue]:
        """根据表头名称读取整列数据。"""

        # load_workbook(...)：打开 Excel 文件，data_only=True 时读取公式的计算结果。
        workbook = self.load_workbook(file_path, data_only=data_only)
        # get_sheet(...)：拿到指定工作表；sheet_name 为 None 时拿当前活动工作表。
        worksheet = self.get_sheet(workbook, sheet_name)
        # read_headers(...)：读取表头行，并把空表头、重复表头做基础处理。
        headers = self.read_headers(worksheet, header_row=header_row)

        column_index = self.find_header_index(headers, column_name)
        values: list[ExcelValue] = []

        # iter_rows(...)：逐行读取数据区；values_only=True 表示只取单元格的值。
        for row in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
            value: ExcelValue = None
            if column_index < len(row):
                # normalize_cell_value(...)：把 Excel 里的值转成业务层常用类型。
                value = self.normalize_cell_value(row[column_index])
            values.append(value)

        return values

    def clean_product_text(self, value: ExcelValue) -> str:
        """清洗商品字段文本，只处理格式问题，不处理业务语义。"""

        if value is None:
            return ""

        # str(...)：把数字、日期等值转成字符串，方便后续统一清洗。
        text = str(value)
        # unicodedata.normalize(...)：把全角字符尽量转成半角字符，例如全角数字转半角数字。
        text = unicodedata.normalize("NFKC", text)
        # replace(...)：把换行、制表符等格式字符替换成普通空格，避免粘在一起。
        text = text.replace("\r", " ")
        text = text.replace("\n", " ")
        text = text.replace("\t", " ")
        text = text.replace("　", " ")
        text = text.replace("\xa0", " ")
        # strip(...)：去掉首尾空格，不删除中间可能有业务意义的内容。
        text = text.strip()
        return text

    def unique_product_values(self, values: Iterable[ExcelValue]) -> list[str]:
        """对商品字段去空、清洗、去重，并保留第一次出现的顺序。"""

        result: list[str] = []
        seen: set[str] = set()

        for value in values:
            cleaned_value = self.clean_product_text(value)
            if not cleaned_value:
                continue
            if cleaned_value in seen:
                continue
            # add(...)：把已经出现过的值记录下来，后面再遇到就跳过。
            seen.add(cleaned_value)
            # append(...)：按原顺序把第一次出现的商品值加入结果列表。
            result.append(cleaned_value)

        return result

    def get_unique_column_values(
        self,
        file_path: str | Path,
        column_name: str,
        *,
        sheet_name: str | None = None,
        header_row: int = 1,
        data_only: bool = True,
    ) -> list[str]:
        """读取指定商品列，并完成基础清洗和去重。"""

        values = self.get_column_values(
            file_path,
            column_name,
            sheet_name=sheet_name,
            header_row=header_row,
            data_only=data_only,
        )
        return self.unique_product_values(values)

    def find_header_index(self, headers: Sequence[str], column_name: str) -> int:
        """在表头列表中查找指定字段的位置。"""

        # strip(...)：去掉传入字段名首尾空格，避免用户输入多余空格导致找不到。
        expected_name = column_name.strip()

        for index, header in enumerate(headers):
            if header == expected_name:
                return index

        available = ", ".join(header for header in headers if header)
        raise ValueError(f"找不到字段：{column_name}；可用字段：{available}")

    def read_headers(self, worksheet: Worksheet, *, header_row: int = 1) -> list[str]:
        """读取并规范化表头。"""

        if header_row < 1:
            raise ValueError("header_row 必须大于等于 1")

        raw_headers = [cell.value for cell in worksheet[header_row]]
        headers: list[str] = []
        seen: dict[str, int] = {}

        for value in raw_headers:
            header = "" if value is None else str(value).strip()
            if not header:
                headers.append("")
                continue

            seen[header] = seen.get(header, 0) + 1
            if seen[header] > 1:
                header = f"{header}_{seen[header]}"
            headers.append(header)

        return self.trim_trailing_empty_headers(headers)

    def write_records(
        self,
        workbook: Workbook,
        sheet_name: str,
        records: Iterable[Record],
        *,
        headers: Sequence[str] | None = None,
        replace_sheet: bool = True,
        auto_width: bool = True,
        freeze_header: bool = True,
        enable_filter: bool = True,
    ) -> Worksheet:
        """把字典列表写入工作表。"""

        materialized_records = list(records)
        resolved_headers = list(headers) if headers is not None else self.infer_headers(materialized_records)
        rows = [[record.get(header) for header in resolved_headers] for record in materialized_records]

        worksheet = self.create_or_replace_sheet(workbook, sheet_name, replace=replace_sheet)
        self.write_table(
            worksheet,
            headers=resolved_headers,
            rows=rows,
            auto_width=auto_width,
            freeze_header=freeze_header,
            enable_filter=enable_filter,
        )
        return worksheet

    def write_table(
        self,
        worksheet: Worksheet,
        *,
        headers: Sequence[str],
        rows: Iterable[RowValues],
        start_row: int = 1,
        start_col: int = 1,
        auto_width: bool = True,
        freeze_header: bool = True,
        enable_filter: bool = True,
    ) -> None:
        """把表头和二维行数据写入工作表。"""

        if start_row < 1 or start_col < 1:
            raise ValueError("start_row 和 start_col 必须大于等于 1")

        for col_offset, header in enumerate(headers):
            worksheet.cell(row=start_row, column=start_col + col_offset, value=header)

        for row_offset, row_values in enumerate(rows, start=1):
            for col_offset, value in enumerate(row_values):
                worksheet.cell(
                    row=start_row + row_offset,
                    column=start_col + col_offset,
                    value=value,
                )

        if headers:
            self.style_header(worksheet, row=start_row, start_col=start_col, end_col=start_col + len(headers) - 1)

            if freeze_header:
                worksheet.freeze_panes = worksheet.cell(row=start_row + 1, column=start_col).coordinate

            if enable_filter:
                end_row = max(worksheet.max_row, start_row)
                end_col_letter = get_column_letter(start_col + len(headers) - 1)
                start_col_letter = get_column_letter(start_col)
                worksheet.auto_filter.ref = f"{start_col_letter}{start_row}:{end_col_letter}{end_row}"

        if auto_width:
            self.auto_fit_columns(worksheet)

    def create_or_replace_sheet(
        self,
        workbook: Workbook,
        sheet_name: str,
        *,
        replace: bool = True,
    ) -> Worksheet:
        """创建工作表；如果同名表已存在，可选择替换或复用。"""

        if sheet_name in workbook.sheetnames:
            if replace:
                old_sheet = workbook[sheet_name]
                index = workbook.index(old_sheet)
                workbook.remove(old_sheet)
                return workbook.create_sheet(sheet_name, index)
            return workbook[sheet_name]

        return workbook.create_sheet(sheet_name)

    def style_header(
        self,
        worksheet: Worksheet,
        *,
        row: int = 1,
        start_col: int = 1,
        end_col: int | None = None,
        fill_color: str = DEFAULT_HEADER_FILL,
    ) -> None:
        """设置表头样式。"""

        end_col = end_col or worksheet.max_column
        fill = PatternFill("solid", fgColor=fill_color)
        font = Font(bold=True, color=self.DEFAULT_HEADER_FONT_COLOR)
        alignment = Alignment(horizontal="center", vertical="center")

        for col in range(start_col, end_col + 1):
            cell = worksheet.cell(row=row, column=col)
            cell.fill = fill
            cell.font = font
            cell.alignment = alignment

    def auto_fit_columns(
        self,
        worksheet: Worksheet,
        *,
        min_width: int = 8,
        max_width: int = 40,
        padding: int = 2,
    ) -> None:
        """根据单元格内容粗略调整列宽。"""

        for column_cells in worksheet.columns:
            column_letter = get_column_letter(column_cells[0].column)
            max_length = 0

            for cell in column_cells:
                if cell.value is None:
                    continue
                max_length = max(max_length, self.display_width(str(cell.value)))

            worksheet.column_dimensions[column_letter].width = min(
                max(max_length + padding, min_width),
                max_width,
            )

    def infer_headers(self, records: Sequence[Record]) -> list[str]:
        """从字典列表中按首次出现顺序推断表头。"""

        headers: list[str] = []
        seen: set[str] = set()
        for record in records:
            for key in record.keys():
                if key not in seen:
                    seen.add(key)
                    headers.append(key)
        return headers

    def trim_trailing_empty_headers(self, headers: Sequence[str]) -> list[str]:
        """移除末尾连续的空表头。"""

        result = list(headers)
        while result and not result[-1]:
            result.pop()
        return result

    def build_preview_cell(self, cell: Any) -> ExcelPreviewCell:
        """构建单元格预览信息，保留 LLM 判断表结构需要的类型和格式。"""

        value = self.normalize_cell_value(cell.value)
        # getattr(...)：安全读取单元格属性；合并单元格 MergedCell 可能没有 is_date 等属性。
        excel_data_type = getattr(cell, "data_type", "")
        # getattr(...)：安全读取数字格式；如果没有 number_format，就按空字符串处理。
        number_format = getattr(cell, "number_format", "")
        # getattr(...)：安全读取日期标记；合并单元格没有 is_date 时默认不是日期。
        is_date = bool(getattr(cell, "is_date", False))
        return ExcelPreviewCell(
            coordinate=cell.coordinate,
            column=cell.column,
            value=value,
            python_type=type(value).__name__ if value is not None else "None",
            excel_data_type=excel_data_type,
            number_format=number_format,
            is_date=is_date,
            display_text=self.format_preview_display_text(value, number_format),
        )

    def trim_trailing_empty_cells(self, cells: Sequence[ExcelPreviewCell]) -> list[ExcelPreviewCell]:
        """移除末尾连续的空单元格预览，让 LLM 输入更紧凑。"""

        result = list(cells)
        while result and (result[-1].value is None or str(result[-1].value).strip() == ""):
            result.pop()
        return result

    def trim_trailing_empty_values(self, values: Sequence[ExcelValue]) -> list[ExcelValue]:
        """移除末尾连续的空单元格值，让预览结果更紧凑。"""

        result = list(values)
        while result and (result[-1] is None or str(result[-1]).strip() == ""):
            result.pop()
        return result

    def format_preview_display_text(self, value: ExcelValue, number_format: str) -> str:
        """把单元格值格式化成更接近 Excel 显示效果的字符串。"""

        if value is None:
            return ""
        if isinstance(value, bool):
            return "TRUE" if value else "FALSE"
        if isinstance(value, datetime):
            return value.isoformat(sep=" ")
        if isinstance(value, date):
            return value.isoformat()
        if isinstance(value, time):
            return value.strftime("%H:%M:%S")
        if isinstance(value, float):
            if "0%" in number_format or "%" in number_format:
                try:
                    return f"{value:.2%}"
                except ValueError:
                    return str(value)
            if number_format in {"General", "@"}:
                return str(value)
            if value.is_integer():
                return str(int(value))
        return str(value)

    def is_empty_row(self, values: Iterable[Any]) -> bool:
        """判断一行是否为空。"""

        return all(value is None or str(value).strip() == "" for value in values)

    def normalize_cell_value(self, value: Any) -> ExcelValue:
        """把 openpyxl 读取的单元格值规整为业务层常用类型，并保留日期/时间类型。"""

        if value is None or isinstance(value, str | int | float | bool | date | datetime | time):
            return value
        return str(value)

    def display_width(self, value: str) -> int:
        """估算字符串在 Excel 中的显示宽度，中文字符按 2 个宽度计算。"""

        return sum(2 if ord(char) > 127 else 1 for char in value)

    def _ensure_xlsx_path(self, file_path: str | Path, *, must_exist: bool) -> Path:
        path = Path(file_path)
        if path.suffix.lower() != ".xlsx":
            raise ValueError(f"仅支持 .xlsx 文件：{path}")
        if must_exist and not path.exists():
            raise FileNotFoundError(f"Excel 文件不存在：{path}")
        return path


excel_tool = ExcelTool()
